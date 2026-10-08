"""Conditional capacity trials remain explicit across previews and seed retries."""
from copy import deepcopy
from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

from test_class_policy_tools import exception, inputs
from test_timetable_agent_tools import write_json
from scripts.timetable_agent_tools import TOOL_SCHEMAS, dispatch_tool, _validate


def assumption(cid='TARGET', expected=50, capacity=30):
    return {'course_id': cid, 'operation': 'set_capacity_assumption',
            'expected_capacity': expected, 'capacity': capacity, 'assumption_only': True,
            'evidence': 'Explicit synthetic trial using an unconfirmed enrollment count'}


def capacity_inputs(root):
    paths = inputs(root, target_flag=1)
    book = load_workbook(paths['course_path'])
    for row in book.active.iter_rows(min_row=3):
        if row[0].value == 'TARGET':
            row[4].value = 50
    book.save(paths['course_path'])
    book.close()
    return paths


def source_hashes(root):
    return {path.name: sha256(path.read_bytes()).hexdigest() for path in root.glob('*.xlsx')}


class CapacityAssumptionToolTests(unittest.TestCase):
    def assert_ok(self, name, arguments):
        schema = next(s['parameters'] for s in TOOL_SCHEMAS if s['name'] == name)
        _validate(arguments, schema)
        response = dispatch_tool(name, arguments)
        self.assertTrue(response['ok'], response)
        self.assertTrue(response['result']['source_files_unchanged'])
        return response

    def assert_assumptions(self, metadata, actions):
        expected = [{key: action[key] for key in
                     ('course_id', 'expected_capacity', 'capacity', 'evidence')}
                    for action in actions if action['operation'] == 'set_capacity_assumption']
        self.assertEqual(metadata['conditional_only'], bool(expected))
        self.assertEqual(metadata['requires_enrollment_confirmation'], bool(expected))
        self.assertEqual(metadata['capacity_course_ids'], sorted(a['course_id'] for a in expected))
        self.assertEqual(metadata['capacity_assumptions'], expected)
        self.assertIsInstance(metadata['interpretation'], str)
        self.assertTrue(metadata['interpretation'].strip())

    def seed(self, root):
        paths = capacity_inputs(root)
        config = {'allowed_changes': [], 'time_scope': 'configured_days', 'days': [0, 1],
                  'periods': [1, 2], 'max_moved_courses': 0,
                  'max_candidates_per_course': 100, 'max_search_nodes': 10000,
                  'time_limit_seconds': 5}
        prefix = [{'course_id': 'GOOD', 'operation': 'set_capacity', 'expected_capacity': 30,
                   'capacity': 35, 'enrollment_frozen': True,
                   'evidence': 'Synthetic confirmed frozen enrollment for GOOD'}]
        seed = self.assert_ok('repair_with_fallbacks', {**paths, 'config': {**config, 'days': [0]},
            'corrections': prefix, 'stages': ['daytime'], 'total_time_limit_seconds': 5})
        proposal = seed['result']['proposal']
        self.assertEqual(proposal['summary']['inserted'], 1)
        self.assertEqual(next(row['status'] for row in proposal['results']
                              if row['course_id'] == 'TARGET'), 'no_matching_room')
        request = {**paths, 'seed_proposal_path': write_json(root / 'seed.json', seed),
                   'target_ids': ['TARGET'], 'corrections': prefix + [assumption()],
                   'config': config, 'stages': ['daytime'], 'total_time_limit_seconds': 5}
        return seed, request

    def test_schema_requires_explicit_assumption_and_positive_capacity(self):
        schema = next(s['parameters'] for s in TOOL_SCHEMAS
                      if s['name'] == 'preview_data_corrections')
        paths = {key: 'synthetic.xlsx' for key in
                 ('course_path', 'room_path', 'schedule_path', 'failed_path')}
        _validate({**paths, 'corrections': [assumption()]}, schema)
        invalid = []
        for key in ('assumption_only', 'expected_capacity', 'capacity', 'evidence'):
            action = assumption(); del action[key]; invalid.append(action)
        for changed in ({'assumption_only': False}, {'assumption_only': 1},
                        {'assumption_only': 'true'}, {'capacity': 0}, {'capacity': -1},
                        {'capacity': True}, {'expected_capacity': -1},
                        {'expected_capacity': True}, {'evidence': ''},
                        {'enrollment_frozen': True}):
            invalid.append({**assumption(), **changed})
        for action in invalid:
            with self.subTest(action=action):
                response = dispatch_tool('preview_data_corrections',
                                         {**paths, 'corrections': [action]})
                self.assertFalse(response['ok'])
                self.assertEqual(response['error']['code'], 'invalid_arguments')

    def test_preview_discloses_unconfirmed_trial_without_editing_source(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); paths = capacity_inputs(root)
            request = {**paths, 'corrections': [assumption()]}
            original = deepcopy(request); before = source_hashes(root)
            response = self.assert_ok('preview_data_corrections', request)
            result = response['result']
            self.assert_assumptions(result['planning_assumptions'], request['corrections'])
            self.assertEqual(result['audit'][0]['operation'], 'set_capacity_assumption')
            self.assertEqual(result['audit'][0]['before']['capacity'], 50)
            self.assertEqual(result['audit'][0]['after']['capacity'], 30)
            self.assertEqual(request, original)
            self.assertEqual(source_hashes(root), before)

    def test_confirmed_capacity_still_requires_frozen_enrollment(self):
        with TemporaryDirectory() as tmp:
            paths = capacity_inputs(Path(tmp))
            action = {'course_id': 'TARGET', 'operation': 'set_capacity',
                      'expected_capacity': 50, 'capacity': 30,
                      'enrollment_frozen': False, 'evidence': 'Unconfirmed count is not frozen'}
            response = dispatch_tool('preview_data_corrections', {**paths, 'corrections': [action]})
            self.assertFalse(response['ok'])
            self.assertEqual(response['error']['code'], 'invalid_arguments')
            action.update(enrollment_frozen=True, evidence='Synthetic confirmed frozen enrollment')
            confirmed = self.assert_ok('preview_data_corrections', {**paths, 'corrections': [action]})
            self.assert_assumptions(confirmed['result']['planning_assumptions'], [action])

    def test_retry_preserves_seed_successes_and_saves_conditional_metadata(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); seed, request = self.seed(root)
            before = source_hashes(root); original = deepcopy(request)
            self.assert_assumptions(seed['result']['proposal']['planning_assumptions'], [])
            response = self.assert_ok('retry_repair_proposal', request)
            proposal = response['result']['proposal']; old = seed['result']['proposal']
            self.assertEqual(proposal['summary']['requested'], 3)
            self.assertEqual(proposal['summary']['inserted'], 2)
            self.assertEqual(proposal['summary']['moved'], 0)
            self.assertEqual(proposal['retry']['newly_inserted'], 1)
            self.assertEqual(proposal['correction_parameters'], request['corrections'])
            self.assert_assumptions(proposal['planning_assumptions'], request['corrections'])
            self.assertEqual(response['result']['planning_assumptions'], proposal['planning_assumptions'])
            for cid in ('BASE', 'GOOD'):
                self.assertEqual([p for p in proposal['placements'] if p['course_id'] == cid],
                                 [p for p in old['placements'] if p['course_id'] == cid])
            self.assertEqual(next(r for r in proposal['results'] if r['course_id'] == 'OTHER'),
                             next(r for r in old['results'] if r['course_id'] == 'OTHER'))
            trial_path = write_json(root / 'trial.json', response)
            diagnosis = dispatch_tool('diagnose_remaining', {'proposal_path': trial_path})
            self.assertTrue(diagnosis['ok'], diagnosis)
            self.assertEqual(diagnosis['result']['planning_assumptions'], proposal['planning_assumptions'])
            self.assertEqual(request, original)
            self.assertEqual(source_hashes(root), before)

    def test_retry_of_other_failure_inherits_seed_capacity_assumption(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); _, request = self.seed(root)
            first = self.assert_ok('retry_repair_proposal', request)
            first_proposal = first['result']['proposal']
            second_request = {**request, 'seed_proposal_path': write_json(root / 'trial.json', first),
                              'target_ids': ['OTHER'],
                              'corrections': request['corrections'] + [exception('OTHER')]}
            second = self.assert_ok('retry_repair_proposal', second_request)
            proposal = second['result']['proposal']
            self.assertEqual(proposal['summary']['requested'], 3)
            self.assertEqual(proposal['summary']['inserted'], 3)
            self.assertEqual(proposal['summary']['moved'], 0)
            self.assertEqual(proposal['retry']['retry_targets'], ['OTHER'])
            self.assertEqual(proposal['planning_assumptions'], first_proposal['planning_assumptions'])
            self.assert_assumptions(proposal['planning_assumptions'], second_request['corrections'])
            for cid in ('BASE', 'GOOD', 'TARGET'):
                self.assertEqual([p for p in proposal['placements'] if p['course_id'] == cid],
                                 [p for p in first_proposal['placements'] if p['course_id'] == cid])
            second_path = write_json(root / 'second.json', second)
            diagnosis = dispatch_tool('diagnose_remaining', {'proposal_path': second_path})
            self.assertTrue(diagnosis['ok'], diagnosis)
            self.assertEqual(diagnosis['result']['planning_assumptions'], first_proposal['planning_assumptions'])

    def test_diagnosis_reconstructs_assumptions_when_cached_marker_is_missing_or_false(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); _, request = self.seed(root)
            trial = self.assert_ok('retry_repair_proposal', request)
            for missing in (True, False):
                saved = deepcopy(trial)
                proposal = saved['result']['proposal']
                if missing:
                    del proposal['planning_assumptions']
                else:
                    proposal['planning_assumptions'] = {
                        'conditional_only': False, 'requires_enrollment_confirmation': False,
                        'capacity_course_ids': [], 'capacity_assumptions': [],
                        'interpretation': 'Incorrect cached claim'}
                with self.subTest(missing=missing):
                    path = write_json(root / 'stale-marker.json', saved)
                    response = dispatch_tool('diagnose_remaining', {'proposal_path': path})
                    self.assertTrue(response['ok'], response)
                    self.assert_assumptions(response['result']['planning_assumptions'], request['corrections'])

    def test_comparison_preserves_scenario_assumptions_without_trusting_cached_markers(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp); _, request = self.seed(root)
            first = self.assert_ok('retry_repair_proposal', request)
            second_request = deepcopy(request)
            second_request['config']['max_search_nodes'] = 20000
            second = self.assert_ok('retry_repair_proposal', second_request)
            self.assertEqual(first['result']['proposal']['source_hashes'],
                             second['result']['proposal']['source_hashes'])
            del first['result']['planning_assumptions']
            del first['result']['proposal']['planning_assumptions']
            false_marker = {'conditional_only': False, 'requires_enrollment_confirmation': False,
                            'capacity_course_ids': [], 'capacity_assumptions': [],
                            'interpretation': 'Incorrect cached claim'}
            second['result']['planning_assumptions'] = deepcopy(false_marker)
            second['result']['proposal']['planning_assumptions'] = deepcopy(false_marker)
            paths = [write_json(root / 'first-scenario.json', first),
                     write_json(root / 'second-scenario.json', second)]
            for proposal_paths in (paths, list(reversed(paths))):
                with self.subTest(proposal_paths=proposal_paths):
                    response = dispatch_tool('compare_proposals', {'proposal_paths': proposal_paths})
                    self.assertTrue(response['ok'], response)
                    result = response['result']
                    self.assert_assumptions(result['planning_assumptions'], request['corrections'])
                    self.assertEqual(result['target_ids'], ['GOOD', 'OTHER', 'TARGET'])
                    self.assertEqual(len(result['proposals']), 2)
                    self.assertIn('max_search_nodes', result['policy_or_budget_differences'])


if __name__ == '__main__':
    unittest.main()
