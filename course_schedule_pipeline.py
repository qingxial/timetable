from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import openpyxl
import xlrd
from openpyxl import Workbook


DISPLAY_HEADERS = [
    "学年学期",
    "课程号",
    "课程名",
    "开课单位代码",
    "开课单位名称",
    "课程类别",
    "上课校区",
    "教学班ID",
    "课序号",
    "课容量",
    "停开课容量",
    "学时",
    "授课学时",
    "上课周次代码",
    "上课周次",
    "课内周学时",
    "教师编号",
    "教师姓名",
    "教师是否通过（1，是，0，否）",
    "是否有主讲资格（1，是，0，否）",
    "账号是否启用（1，是，0，否）",
    "教师上课周次代码",
    "教师上课周次名称",
    "教师职称代码",
    "教师职称名称",
    "课程主讲教师代码",
    "课程主讲教师名称",
    "是否需要排课（1，是，0，否）",
    "是否需要教室（1，是，0，否）",
    "教室类型代码",
    "教室类型名称",
    "教学楼代码，名称见教室资源表",
    "教室代码，名称见教室资源表",
    "偏好上课时间",
    "避免排课时间",
    "排课要求描述",
    "历史教室代码",
    "推荐班级",
    None,
    None,
]

CODE_HEADERS = [
    "XNXQDM",
    "KCH",
    "KCM",
    "KKDWDM",
    "YXMC",
    "KCLB",
    "SKXQ",
    "JXBID",
    "KXH",
    "KRL",
    "TKKRL",
    "XS",
    "SKXS",
    "SKZCDM",
    "SKZCMC",
    "ZXS",
    "JSH",
    "XM",
    "SFTGPX",
    "SFYZJZG",
    "ZHZT",
    "RWJSZCDM",
    "RWJSZCMC",
    "ZCDM",
    "ZCMC",
    "SFFZJSDM",
    "SFZJJSMC",
    "SFXYPK",
    "SFXYJAS",
    "JASLXDM",
    "JASLXMC",
    "JXLDM",
    "JASDM",
    "Prefer_Time",
    "unavailable_Time",
    "PKYQMS",
    "LSJASDM",
    "TJBJ",
    None,
    None,
]


def clean_text(value: Any) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.lower() == "nan":
        return None
    if text.endswith(".0") and re.fullmatch(r"\d+\.0", text):
        return text[:-2]
    return text


def normalize_key_value(value: Any) -> str:
    return clean_text(value) or ""


def yes_no_to_flag(value: Any, default: str = "1") -> str:
    text = clean_text(value)
    if text is None:
        return default
    if text in {"是", "Y", "y", "Yes", "YES", "1", "true", "True"}:
        return "1"
    if text in {"否", "N", "n", "No", "NO", "0", "false", "False"}:
        return "0"
    return default


def split_names(value: Any) -> List[str]:
    text = clean_text(value)
    if not text:
        return []
    return [part.strip() for part in re.split(r"[,，、;；]", text) if part.strip()]


def parse_week_code(week_text: Any, total_weeks: int = 16) -> str:
    text = clean_text(week_text)
    bits = ["0"] * total_weeks
    if not text:
        return "".join(bits)

    for part in re.split(r"[,，;；、\s]+", text):
        part = part.strip()
        if not part:
            continue
        odd_only = "单" in part
        even_only = "双" in part
        for start_text, end_text in re.findall(r"(\d+)(?:\s*-\s*(\d+))?", part):
            start = int(start_text)
            end = int(end_text or start_text)
            if start > end:
                start, end = end, start
            for week in range(start, end + 1):
                if week < 1 or week > total_weeks:
                    continue
                if odd_only and week % 2 == 0:
                    continue
                if even_only and week % 2 == 1:
                    continue
                bits[week - 1] = "1"
    return "".join(bits)


def read_table(path: Path) -> List[Dict[str, Any]]:
    suffix = path.suffix.lower()
    if suffix == ".xls":
        return read_xls_table(path)
    if suffix == ".xlsx":
        return read_xlsx_table(path)
    raise ValueError(f"Unsupported Excel extension: {path}")


def read_xls_table(path: Path) -> List[Dict[str, Any]]:
    workbook = xlrd.open_workbook(str(path))
    #workbook=Workbook(str(path))
    sheet = workbook.sheet_by_index(0)
    headers = [clean_text(sheet.cell_value(0, col)) for col in range(sheet.ncols)]
    rows: List[Dict[str, Any]] = []
    for row_idx in range(1, sheet.nrows):
        row: Dict[str, Any] = {}
        has_value = False
        for col_idx, header in enumerate(headers):
            if not header:
                continue
            value = sheet.cell_value(row_idx, col_idx)
            value = None if value == "" else value
            if value is not None:
                has_value = True
            row[header] = value
        if has_value:
            rows.append(row)
    return rows


def read_xlsx_table(path: Path, skip_code_row: bool = False) -> List[Dict[str, Any]]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows_iter = worksheet.iter_rows(values_only=True)
    headers = [clean_text(value) for value in next(rows_iter)]
    if skip_code_row:
        next(rows_iter, None)

    rows: List[Dict[str, Any]] = []
    for values in rows_iter:
        row: Dict[str, Any] = {}
        has_value = False
        for header, value in zip(headers, values):
            if not header:
                continue
            if value is not None:
                has_value = True
            row[header] = value
        if has_value:
            rows.append(row)
    return rows


@dataclass
class BuildStats:
    teaching_rows: int = 0
    output_rows: int = 0
    missing_teacher_names: List[str] = field(default_factory=list)
    duplicate_enrollment_keys: int = 0
    missing_enrollment_keys: int = 0
    missing_classroom_keys: int = 0


@dataclass
class CourseSchedulePipeline:
    teaching_task_path: Path
    enrollment_course_path: Path
    output_path: Path
    semester: str = "2025-2026-1"
    teacher_roster_path: Optional[Path] = None
    classroom_history_path: Optional[Path] = None
    total_weeks: int = 16

    def run(self) -> BuildStats:
        teaching_rows = read_table(self.teaching_task_path)
        enrollment_rows = read_table(self.enrollment_course_path)
        teacher_map = self.load_teacher_map()
        classroom_map = self.load_classroom_history_map()
        enrollment_map, duplicate_enrollment_keys = self.build_enrollment_map(enrollment_rows)

        stats = BuildStats(
            teaching_rows=len(teaching_rows),
            duplicate_enrollment_keys=duplicate_enrollment_keys,
        )
        output_rows: List[List[Any]] = []
        missing_teacher_names: set[str] = set()
        missing_enrollment_keys: set[Tuple[str, str]] = set()
        missing_classroom_keys: set[Tuple[str, str]] = set()

        next_new_id = self.next_new_teacher_id(teacher_map.values())
        for source_row in teaching_rows:
            key = self.course_key(source_row)
            enrollment_row = enrollment_map.get(key)
            if enrollment_row is None:
                missing_enrollment_keys.add(key)

            campus = source_row.get("学校校区")
            classroom_key = (normalize_key_value(source_row.get("课程号")), normalize_key_value(campus))
            history_classroom_code = classroom_map.get(classroom_key)
            if self.classroom_history_path and history_classroom_code is None:
                missing_classroom_keys.add(classroom_key)

            week_name = source_row.get("上课周次")
            week_code = parse_week_code(week_name, self.total_weeks)
            teachers = split_names(source_row.get("上课教师")) or [None]

            for teacher_name in teachers:
                teacher_id = None
                if teacher_name:
                    teacher_id = teacher_map.get(teacher_name)
                    if teacher_id is None:
                        teacher_id = f"new{next_new_id}"
                        next_new_id += 1
                        teacher_map[teacher_name] = teacher_id
                        missing_teacher_names.add(teacher_name)

                output_rows.append(
                    self.make_output_row(
                        source_row=source_row,
                        enrollment_row=enrollment_row,
                        teacher_name=teacher_name,
                        teacher_id=teacher_id,
                        week_code=week_code,
                        history_classroom_code=history_classroom_code,
                    )
                )

        self.write_output(output_rows)
        stats.output_rows = len(output_rows)
        stats.missing_teacher_names = sorted(missing_teacher_names)
        stats.missing_enrollment_keys = len(missing_enrollment_keys)
        stats.missing_classroom_keys = len(missing_classroom_keys)
        return stats

    def load_teacher_map(self) -> Dict[str, str]:
        if not self.teacher_roster_path:
            return {}
        rows = read_xlsx_table(self.teacher_roster_path)
        teacher_map: Dict[str, str] = {}
        for row in rows:
            teacher_id = clean_text(row.get("JSH"))
            teacher_name = clean_text(row.get("XM"))
            if teacher_id and teacher_name and teacher_name not in teacher_map:
                teacher_map[teacher_name] = teacher_id
        return teacher_map

    def load_classroom_history_map(self) -> Dict[Tuple[str, str], str]:
        if not self.classroom_history_path:
            return {}
        rows = read_xlsx_table(self.classroom_history_path)
        classroom_map: Dict[Tuple[str, str], str] = {}
        for row in rows:
            key = (normalize_key_value(row.get("课程号")), normalize_key_value(row.get("学校校区")))
            code = clean_text(row.get("历史教室代码"))
            if key[0] and key[1] and code and key not in classroom_map:
                classroom_map[key] = code
        return classroom_map

    @staticmethod
    def build_enrollment_map(rows: Sequence[Dict[str, Any]]) -> Tuple[Dict[Tuple[str, str], Dict[str, Any]], int]:
        result: Dict[Tuple[str, str], Dict[str, Any]] = {}
        duplicate_count = 0
        for row in rows:
            key = (normalize_key_value(row.get("课程号")), normalize_key_value(row.get("课序号")))
            if not key[0] or not key[1]:
                continue
            if key in result:
                duplicate_count += 1
                continue
            result[key] = row
        return result, duplicate_count

    @staticmethod
    def course_key(row: Dict[str, Any]) -> Tuple[str, str]:
        return (normalize_key_value(row.get("课程号")), normalize_key_value(row.get("课序号")))

    @staticmethod
    def next_new_teacher_id(existing_ids: Iterable[str]) -> int:
        max_id = 0
        for teacher_id in existing_ids:
            match = re.fullmatch(r"new(\d+)", teacher_id)
            if match:
                max_id = max(max_id, int(match.group(1)))
        return max_id + 1

    def make_output_row(
        self,
        source_row: Dict[str, Any],
        enrollment_row: Optional[Dict[str, Any]],
        teacher_name: Optional[str],
        teacher_id: Optional[str],
        week_code: str,
        history_classroom_code: Optional[str],
    ) -> List[Any]:
        zxs = enrollment_row.get("课内周学时") if enrollment_row else None
        recommended_class = enrollment_row.get("推荐班级") if enrollment_row else source_row.get("上课班级")
        return [
            self.semester,
            source_row.get("课程号"),
            source_row.get("课程名"),
            None,
            source_row.get("开课单位"),
            source_row.get("课程类别"),
            source_row.get("学校校区"),
            source_row.get("教学班ID"),
            source_row.get("课序号"),
            source_row.get("课容量"),
            source_row.get("停开课容量"),
            source_row.get("学时"),
            source_row.get("授课学时"),
            week_code,
            source_row.get("上课周次"),
            zxs,
            teacher_id,
            teacher_name,
            None,
            None,
            None,
            week_code,
            None,
            None,
            None,
            None,
            None,
            yes_no_to_flag(source_row.get("是否需要排课"), default="1"),
            yes_no_to_flag(source_row.get("是否需要教室"), default="1"),
            None,
            None,
            None,
            None,
            None,
            None,
            source_row.get("排课要求描述") or "",
            history_classroom_code,
            recommended_class,
            None,
            None,
        ]

    def write_output(self, rows: Sequence[Sequence[Any]]) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Sheet1"
        worksheet.append(DISPLAY_HEADERS)
        worksheet.append(CODE_HEADERS)
        for row in rows:
            worksheet.append(list(row))
        workbook.save(self.output_path)


def default_data_dir() -> Path:
    return Path(__file__).resolve().parent / "智能排课基础数据"


def build_default_pipeline(output_path: Optional[Path] = None) -> CourseSchedulePipeline:
    data_dir = default_data_dir()
    return CourseSchedulePipeline(
        teaching_task_path=data_dir / "2025.10.20-2025-2026-1教学任务.xls",
        enrollment_course_path=data_dir / "2025.10.20-2025-2026-1选课课程- (1).xls",
        output_path=output_path or data_dir / "课程表2025-2026-1_pipeline.xlsx",
        teacher_roster_path=data_dir / "更新后的教师名单2025-2026-1.xlsx",
        classroom_history_path=data_dir / "课程教室对照表.xlsx",
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build the 2025-2026-1 course schedule resource table.")
    parser.add_argument("--teaching-task", type=Path, help="Path to 教学任务.xls")
    parser.add_argument("--enrollment-course", type=Path, help="Path to 选课课程.xls/xlsx")
    parser.add_argument("--teacher-roster", type=Path, help="Optional teacher roster with JSH/XM columns")
    parser.add_argument("--classroom-history", type=Path, help="Optional 课程教室对照表.xlsx")
    parser.add_argument("--output", type=Path, help="Output xlsx path")
    parser.add_argument("--semester", default="2025-2026-1")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if args.teaching_task and args.enrollment_course and args.output:
        pipeline = CourseSchedulePipeline(
            teaching_task_path=args.teaching_task,
            enrollment_course_path=args.enrollment_course,
            output_path=args.output,
            semester=args.semester,
            teacher_roster_path=args.teacher_roster,
            classroom_history_path=args.classroom_history,
        )
    else:
        pipeline = build_default_pipeline(args.output)

    stats = pipeline.run()
    print(f"teaching_rows={stats.teaching_rows}")
    print(f"output_rows={stats.output_rows}")
    print(f"duplicate_enrollment_keys={stats.duplicate_enrollment_keys}")
    print(f"missing_enrollment_keys={stats.missing_enrollment_keys}")
    print(f"missing_classroom_keys={stats.missing_classroom_keys}")
    print(f"missing_teacher_names={len(stats.missing_teacher_names)}")
    if stats.missing_teacher_names:
        print(",".join(stats.missing_teacher_names[:50]))


if __name__ == "__main__":
    main()
