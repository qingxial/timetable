# -*- coding: utf-8 -*-
"""把课程表里 ZXS 为奇数（3/5/7）的班拆成两条虚班记录：
- 虚班 A: JXBID = 原+'@W'，ZXS = 原-1（偶数），SKZCDM 不变（每周）
- 虚班 B: JXBID = 原+'@B'，ZXS = 2，SKZCDM = 原 SKZCDM ∩ 双周掩码（隔周一次2节）

通式验证：ZXS=3 → 2 每周 + 2 双周 ; ZXS=5 → 4 每周 + 2 双周 ; ZXS=7 → 6 每周 + 2 双周
平均周课时严格守恒：18×(ZXS-1) + 9×2 = 18×ZXS ✓（9 = ⌊18/2⌋ 双周周数）

输入：智能排课基础数据/提取的基础数据表_converted/课程表.xlsx
输出：智能排课基础数据/提取的基础数据表_converted/课程表_split.xlsx
保留 2 行表头；其他字段全部保留原值。

用法：
  .venv/bin/python scripts/split_odd_zxs.py            # dry-run 预览
  .venv/bin/python scripts/split_odd_zxs.py --apply    # 真生成新文件
"""
import argparse, os
from copy import copy
from openpyxl import load_workbook

SRC = "智能排课基础数据/提取的基础数据表_converted/课程表.xlsx"
DST = "智能排课基础数据/提取的基础数据表_converted/课程表_split.xlsx"
TOTAL_WEEKS = 23

# 需要拆分的 ZXS（含小数容差：3.0/3.0000001 等）
TARGET_ZXS = {3, 5, 7}


def _to_int_zxs(v):
    """把 ZXS 转成 int；容错 NaN/字符串/小数。返回 None 表示无法判断。"""
    try:
        f = float(v)
    except (TypeError, ValueError):
        return None
    # 容差 1e-3 处理 3.9999998 这种浮点漂移
    n = round(f)
    if abs(f - n) > 1e-3:
        return None
    return int(n)


def biweekly_mask(skzcdm: str) -> str:
    """根据 SKZCDM 第一个 '1' 位的奇偶，生成「原 SKZCDM ∩ 双周掩码」。

    第一个 '1' 在第 w 周（1-based）：
      - w 奇 → 保留奇数位（3,5,7,... 这种"单周"）
      - w 偶 → 保留偶数位（2,4,6,... 这种"双周"）
    """
    s = str(skzcdm).strip()
    if len(s) != TOTAL_WEEKS or '1' not in s:
        return '0' * TOTAL_WEEKS
    first_one = s.index('1') + 1  # 1-based week
    parity = first_one % 2  # 1=奇, 0=偶
    out = []
    for i, b in enumerate(s):
        week = i + 1
        if b == '1' and (week % 2) == parity:
            out.append('1')
        else:
            out.append('0')
    return ''.join(out)


def bits_to_csv(skzcdm: str) -> str:
    return ','.join(str(i + 1) for i, b in enumerate(skzcdm) if b == '1')


def split_row(row_values, col_idx):
    """row_values: dict {col_name: value}；col_idx: {col_name: 1-based index}。
    返回 (A 行值, B 行值) 或 None（无需拆分）。

    拆分条件（学时严格守恒前提）：
    - ZXS ∈ {3, 5, 7}
    - W = 原 SKZCDM 上课周数 必须 ≥ 4 且为偶数
      （奇数 W 拆「每周+双周」会差 ±1 学时；W<4 太短无拆分意义）
    - 双周掩码后至少 1 周
    """
    zxs = _to_int_zxs(row_values.get("ZXS"))
    if zxs not in TARGET_ZXS:
        return None
    skzcdm = str(row_values.get("SKZCDM") or "").strip()
    if len(skzcdm) != TOTAL_WEEKS or '1' not in skzcdm:
        return None
    W = skzcdm.count('1')
    if W < 4 or W % 2 != 0:
        return None    # W 奇数或 < 4：不拆（保留原 ZXS，走现有 day_pattern）
    jxbid = str(row_values.get("JXBID") or "").strip()
    if not jxbid:
        return None

    # 虚班 B 的 SKZCDM
    skz_b = biweekly_mask(skzcdm)
    if '1' not in skz_b:
        return None  # 双周掩码后全 0，拆分无意义

    # 学时严格守恒校验：A 周数 × (ZXS-1) + B 周数 × 2 == W × ZXS
    b_weeks = skz_b.count('1')
    if W * (zxs - 1) + b_weeks * 2 != W * zxs:
        return None  # 不守恒则不拆（保险）

    # 虚班 A: 原 SKZCDM 不变；ZXS = 原-1
    a = dict(row_values)
    a["JXBID"] = jxbid + "@W"
    a["ZXS"] = zxs - 1
    # SKZCDM / SKZCMC 保持原值

    # 虚班 B: SKZCDM = 双周掩码；ZXS = 2
    b = dict(row_values)
    b["JXBID"] = jxbid + "@B"
    b["ZXS"] = 2
    b["SKZCDM"] = skz_b
    if "SKZCMC" in col_idx:
        b["SKZCMC"] = bits_to_csv(skz_b)

    # 同步教师周次代码（如果它原来等于 SKZCDM，说明教师全程跟课，虚班 B 也要同步收窄）
    if "RWJSZCDM" in col_idx:
        rwj = str(row_values.get("RWJSZCDM") or "").strip()
        if rwj == skzcdm:
            b["RWJSZCDM"] = skz_b
            if "RWJSZCMC" in col_idx:
                b["RWJSZCMC"] = bits_to_csv(skz_b)
    return a, b


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="真生成 DST 文件，否则仅 dry-run")
    args = ap.parse_args()

    wb = load_workbook(SRC)
    ws = wb.active
    header_en = [c.value for c in ws[2]]
    col_idx = {h: i + 1 for i, h in enumerate(header_en) if h}

    required = ["JXBID", "ZXS", "SKZCDM"]
    miss = [c for c in required if c not in col_idx]
    if miss:
        raise RuntimeError(f"缺列: {miss}")

    # 收集所有数据行
    data_rows = []
    for r in range(3, ws.max_row + 1):
        row = {h: ws.cell(r, col_idx[h]).value for h in header_en if h}
        data_rows.append(row)

    split_count = 0
    samples = []
    new_rows = []  # 处理后的全部数据行
    for row in data_rows:
        out = split_row(row, col_idx)
        if out is None:
            new_rows.append(row)
            continue
        a, b = out
        new_rows.append(a)
        new_rows.append(b)
        split_count += 1
        if len(samples) < 6:
            samples.append((row["JXBID"], row.get("KCM"), row["ZXS"], row["SKZCDM"],
                            a["JXBID"], a["ZXS"], a["SKZCDM"],
                            b["JXBID"], b["ZXS"], b["SKZCDM"]))

    print(f"原行数: {len(data_rows)}")
    print(f"拆分行数: {split_count}（→ 增加 {split_count} 行）")
    print(f"新行数: {len(new_rows)}")
    print()
    print("=== 拆分样例（前 6 组）===")
    for s in samples:
        print(f"原 [{s[0]}] {s[1]}  ZXS={s[2]}  SKZCDM={s[3]}")
        print(f"  → A [{s[4]}]  ZXS={s[5]}  SKZCDM={s[6]}")
        print(f"  → B [{s[7]}]  ZXS={s[8]}  SKZCDM={s[9]}")
        # 学时验证：A 实际上 a_W 周 × A.ZXS；B 实际上 b_W 周 × B.ZXS
        orig_W = sum(1 for c in s[3] if c == '1')
        a_W = sum(1 for c in s[6] if c == '1')
        b_W = sum(1 for c in s[9] if c == '1')
        a_zxs_n = int(s[5]); b_zxs_n = int(s[8])
        orig_zxs = int(s[2]) if isinstance(s[2], (int, float)) else int(float(s[2]))
        a_h = a_W * a_zxs_n; b_h = b_W * b_zxs_n
        print(f"     验证：A({a_W}周×{a_zxs_n}={a_h}) + B({b_W}周×{b_zxs_n}={b_h}) = {a_h+b_h}  vs  原 W×ZXS = {orig_W}×{orig_zxs} = {orig_W*orig_zxs}")
        print()

    if not args.apply:
        print("[dry-run] 未写入。加 --apply 真生成新文件。")
        return

    # 用 openpyxl 写新工作簿（复制原表头格式）
    from openpyxl import Workbook
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = ws.title
    # 第 1 行（中文表头）
    for c, v in enumerate([ws.cell(1, i + 1).value for i in range(len(header_en))], start=1):
        ws2.cell(1, c).value = v
    # 第 2 行（英文表头）
    for c, v in enumerate(header_en, start=1):
        ws2.cell(2, c).value = v
    # 数据行
    for ri, row in enumerate(new_rows, start=3):
        for h, ci in col_idx.items():
            ws2.cell(ri, ci).value = row.get(h)
    wb2.save(DST)
    print(f"\n已写出 → {DST}  （{len(new_rows)} 行）")


if __name__ == "__main__":
    main()
