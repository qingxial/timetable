"""Lossless, pickle-free inputs for auditable timetable repair.

Weeks and periods are one-based; days are zero-based. Invalid baseline rows
remain represented (with empty weeks/periods or day=-1), with ``baseline:``
issues. Explicit zero-hour, reversed-period export placeholders are retained
only in diagnostics, because they do not represent teaching or occupancy.
"""
from __future__ import annotations

from dataclasses import dataclass, field, fields, is_dataclass
from hashlib import sha256
from math import isfinite
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook


DAYS = {f"周{char}": index for index, char in enumerate("一二三四五六日")}
PERIODS_PER_DAY = 11


@dataclass
class Course:
    id: str
    name: str = ""
    hours: float = 0.0
    weeks: frozenset[int] = frozenset()
    teachers: dict[str, frozenset[int]] = field(default_factory=dict)
    classes: frozenset[str] = frozenset()
    campus: str = ""
    capacity: float = 0.0
    room_type: str = ""
    buildings: frozenset[str] = frozenset()
    explicit_rooms: frozenset[str] = frozenset()
    historical_rooms: frozenset[str] = frozenset()
    prefer: str = ""
    unavailable: str = ""
    special: str = ""
    max_block: int = 2
    block_template: tuple[int, ...] = ()
    check_room: bool = True
    check_class: bool = True
    issues: list[str] = field(default_factory=list)
    total_hours: float | None = None
    weekly_loads: dict[int, int] = field(default_factory=dict)


@dataclass
class Room:
    id: str
    name: str = ""
    campus: str = ""
    capacity: float = 0.0
    type: str = ""
    building: str = ""
    enabled: bool = False


@dataclass(frozen=True)
class Placement:
    course_id: str
    weeks: frozenset[int]
    day: int
    periods: tuple[int, ...]
    room_id: str


@dataclass
class Dataset:
    courses: dict[str, Course] = field(default_factory=dict)
    rooms: dict[str, Room] = field(default_factory=dict)
    placements: list[Placement] = field(default_factory=list)
    issues: list[str] = field(default_factory=list)
    source_hashes: dict[str, str] = field(default_factory=dict)
    class_registry: frozenset[str] | None = None

    def to_dict(self) -> dict[str, Any]:
        """Return deterministic JSON-compatible data, including diagnostics."""
        return _json_value(self)


def _json_value(value: Any) -> Any:
    if is_dataclass(value):
        return {f.name: _json_value(getattr(value, f.name)) for f in fields(value)}
    if isinstance(value, dict):
        return {str(k): _json_value(v) for k, v in value.items()}
    if isinstance(value, (set, frozenset)):
        return [_json_value(v) for v in sorted(value)]
    if isinstance(value, (list, tuple)):
        return [_json_value(v) for v in value]
    return value


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        if not isfinite(value):
            return ""
        if value.is_integer():
            return str(int(value))
    return str(value).strip()


def _normalize_time(text: str) -> str:
    return text.translate(str.maketrans({
        "（": "(", "）": ")", "；": ";", "，": ",",
        "【": "[", "】": "]", "［": "[", "］": "]",
        "－": "-", "–": "-", "—": "-",
    }))


def parse_time_windows(text: str) -> set[tuple[int, int]]:
    """Parse canonical time windows; never silently accept partial matches.

    Accept 周一(1-2节), 全周(1-2节), 周日全天, and front-end square
    bracket groups.  Semicolons, commas and Chinese list punctuation separate
    complete windows.  ``全天`` means periods 1 through 11 on the named day.
    """
    normalized = _normalize_time(_text(text))
    if not normalized:
        return set()
    # Brackets group the front-end's canonical windows, not arbitrary text.
    depth = 0
    for char in normalized:
        if char == "[":
            depth += 1
        elif char == "]":
            depth -= 1
        if depth < 0:
            raise ValueError(f"unbalanced time-window brackets: {text!r}")
    if depth:
        raise ValueError(f"unbalanced time-window brackets: {text!r}")
    normalized = normalized.replace("[", "").replace("]", "")
    parts = re.split(r"[;、,]", normalized)
    windows: set[tuple[int, int]] = set()
    for index, part in enumerate(parts):
        part = re.sub(r"\s+", "", part)
        if not part:
            # A final separator is harmless; an empty expression is not.
            if index == len(parts) - 1 and windows:
                continue
            raise ValueError(f"empty time window in {text!r}")
        match = re.fullmatch(r"(周[一二三四五六日]|全周)(?:\((\d+)-(\d+)节\)|(全天))", part)
        if not match:
            raise ValueError(f"unsupported time window: {part!r}")
        day_name, start, end, all_day = match.groups()
        a, b = (1, PERIODS_PER_DAY) if all_day else (int(start), int(end))
        if not 1 <= a <= b <= PERIODS_PER_DAY:
            raise ValueError(f"invalid period range: {part!r}")
        days = range(7) if day_name == "全周" else [DAYS[day_name]]
        windows.update((day, period) for day in days for period in range(a, b + 1))
    return windows


def _codes(value: Any) -> frozenset[str]:
    return frozenset(x for x in re.split(r"[,，、;；\s]+", _text(value)) if x)


def _number(value: Any, name: str, issues: list[str], *, default: float | None = None) -> float:
    text = _text(value)
    if not text and default is not None:
        return float(default)
    try:
        number = float(text)
        if not isfinite(number) or number < 0:
            raise ValueError()
        return number
    except (ValueError, TypeError):
        issues.append(f"metadata: invalid {name}: {text!r}")
        return 0.0


def _week_flags(value: Any) -> frozenset[int]:
    text = _text(value)
    if not re.fullmatch(r"[01]{1,53}", text):
        raise ValueError(f"invalid binary week code: {text!r}")
    return frozenset(i + 1 for i, flag in enumerate(text) if flag == "1")


def _week_list(value: Any) -> frozenset[int]:
    text = _text(value)
    if not text:
        raise ValueError("missing teaching weeks")
    weeks: set[int] = set()
    for part in re.split(r"[,，、;；\s]+", text):
        match = re.fullmatch(r"(\d+)(?:[-~至](\d+))?", part)
        if not match:
            raise ValueError(f"invalid teaching weeks: {text!r}")
        a, b = int(match[1]), int(match[2] or match[1])
        if not 1 <= a <= b <= 53:
            raise ValueError(f"invalid teaching week range: {part!r}")
        weeks.update(range(a, b + 1))
    return frozenset(weeks)


def _periods(value: Any) -> tuple[int, ...]:
    text = _normalize_time(_text(value))
    match = re.fullmatch(r"第?(\d+)(?:-(\d+))?节?", text)
    if not match:
        raise ValueError(f"unsupported periods: {text!r}")
    a, b = int(match[1]), int(match[2] or match[1])
    if not 1 <= a <= b <= PERIODS_PER_DAY:
        raise ValueError(f"invalid periods: {text!r}")
    return tuple(range(a, b + 1))


def _check_flag(value: Any, name: str, issues: list[str]) -> bool:
    text = _text(value)
    if text not in ("", "0", "0.0", "1", "1.0"):
        issues.append(f"metadata: invalid {name}: {text!r}")
    # The data owner's semantics override the misleading Chinese headers.
    return text not in ("1", "1.0")


def _read_table(path: Path, double_header: bool) -> list[tuple[int, dict[str, Any]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        chinese = next(rows, ())
        headers = next(rows, ()) if double_header else chinese
        resolved = []
        for index, raw in enumerate(headers):
            header = _text(raw)
            if not header and double_header and index < len(chinese):
                cn = _text(chinese[index])
                for prefix, key in (("理论学时", "LLXS"), ("是否检查教室冲突", "IF_ROOM_CONFICT"),
                                    ("是否检查班级冲突", "IF_CLASS_CONFICT")):
                    if cn.startswith(prefix):
                        header = key
            resolved.append(header)
        result = []
        for number, row in enumerate(rows, start=3 if double_header else 2):
            if any(_text(value) for value in row):
                result.append((number, {key: value for key, value in zip(resolved, row) if key}))
        return result
    finally:
        workbook.close()


def _load_courses(rows: list[tuple[int, dict[str, Any]]], dataset: Dataset,
                  class_registry: frozenset[str] | None = None) -> None:
    grouped: dict[str, list[tuple[int, dict[str, Any]]]] = {}
    for number, row in rows:
        identifier = _text(row.get("JXBID"))
        if not identifier:
            dataset.issues.append(f"metadata:course row {number}: missing JXBID")
            continue
        grouped.setdefault(identifier, []).append((number, row))
    # Week fields and teachers deliberately differ across rows: they are unions,
    # whereas contradictory shared constraints cannot safely select a first row.
    shared = ("KCM", "ZXS", "KRL", "SKXQ", "JASLXMC", "JASLXDM", "JXLDM", "JASDM",
              "LSJASDM", "Prefer_Time", "unavailable_Time", "PKYQMS", "TJBJ",
              "IF_ROOM_CONFICT", "IF_CLASS_CONFICT", "max_block", "block_template",
              "SFXYPK", "SFXYJAS")
    for identifier, records in grouped.items():
        first = records[0][1]
        issues: list[str] = []
        for key in shared:
            values = {_text(row.get(key)) for _, row in records}
            if len(values) > 1:
                issues.append(f"metadata: inconsistent {key} across course rows: {sorted(values)!r}")
        total_values: list[float | None] = []
        for number, row in records:
            raw_total = _text(row.get("LLXS"))
            if not raw_total:
                total_values.append(None)
                continue
            try:
                value = float(raw_total)
                if not isfinite(value) or value < 0:
                    raise ValueError()
                # Zero is not a declared positive theory-hour requirement.
                total_values.append(value if value > 0 else None)
            except ValueError:
                total_values.append(None)
                issues.append(f"metadata: row {number}: invalid LLXS: {raw_total!r}")
        if len(set(total_values)) > 1:
            issues.append(f"metadata: inconsistent LLXS across course rows: {total_values!r}")
            total_hours = None
        else:
            total_hours = total_values[0]
        weeks: set[int] = set()
        teachers: dict[str, set[int]] = {}
        for number, row in records:
            try:
                weeks.update(_week_flags(row.get("SKZCDM")))
            except ValueError as error:
                issues.append(f"metadata: row {number} SKZCDM: {error}")
            teacher = _text(row.get("JSH"))
            if not teacher:
                issues.append(f"metadata: row {number}: missing teacher identifier")
                continue
            try:
                teachers.setdefault(teacher, set()).update(_week_flags(row.get("RWJSZCDM")))
            except ValueError as error:
                teachers.setdefault(teacher, set())
                issues.append(f"metadata: row {number} RWJSZCDM: {error}")
        if not weeks:
            issues.append("metadata: no active teaching weeks")
        teacher_weeks = {teacher: frozenset(active & weeks) for teacher, active in teachers.items()}
        for teacher, active in teacher_weeks.items():
            if not active:
                issues.append(f"metadata: teacher {teacher!r} has no responsibility weeks within teaching weeks")
        # Keep every known class for conservative baseline occupancy even when
        # inconsistent metadata makes this course ineligible for new repair.
        classes = {class_id for _, row in records for class_id in _codes(row.get("TJBJ"))}
        generic = {c for c in classes if c.startswith(("全校", "全院", "全学院"))}
        if generic:
            classes.difference_update(generic)
            issues.append(f"coverage: generic class scopes cannot identify student conflicts: {sorted(generic)!r}")
        if not classes and not generic:
            issues.append("coverage: missing class identifiers")
        check_class = _check_flag(first.get("IF_CLASS_CONFICT"), "IF_CLASS_CONFICT", issues)
        # The registry, when supplied, is authoritative even for real class
        # names without digits. Keep unknown tokens for baseline reservations.
        unregistered = classes if class_registry is None else classes - class_registry
        malformed = sorted(c for c in unregistered if not any(char.isdigit() for char in c) or c.startswith("周"))
        if malformed:
            issues.append(f"coverage: unresolved class identifiers: {malformed!r}")
        if class_registry is not None and check_class and unregistered:
            issues.append(f"coverage: class identifiers absent from BJMC registry: {sorted(unregistered)!r}")
        for key in ("Prefer_Time", "unavailable_Time"):
            for value in {_text(row.get(key)) for _, row in records}:
                try:
                    parse_time_windows(value)
                except ValueError as error:
                    issues.append(f"constraint: {key}: {error}")
        max_block = _number(first.get("max_block"), "max_block", issues, default=2)
        if not max_block.is_integer() or not 1 <= max_block <= PERIODS_PER_DAY:
            issues.append(f"metadata: max_block must be an integer from 1 to {PERIODS_PER_DAY}")
            max_block = 2.0
        template: tuple[int, ...] = ()
        template_text = _text(first.get("block_template"))
        if template_text:
            try:
                pieces = re.split(r"[,，、;；\s]+", template_text)
                template = tuple(int(piece) for piece in pieces)
                if not template or any(not 1 <= item <= PERIODS_PER_DAY for item in template):
                    raise ValueError()
            except ValueError:
                template = ()
                issues.append(f"metadata: invalid block_template: {template_text!r}")
        campus = _text(first.get("SKXQ"))
        if not campus:
            issues.append("metadata: missing campus")
        course = Course(
            id=identifier, name=_text(first.get("KCM")),
            hours=_number(first.get("ZXS"), "ZXS", issues), weeks=frozenset(weeks),
            teachers=teacher_weeks, classes=frozenset(classes), campus=campus,
            capacity=_number(first.get("KRL"), "KRL", issues),
            room_type=_text(first.get("JASLXMC")) or _text(first.get("JASLXDM")),
            buildings=_codes(first.get("JXLDM")), explicit_rooms=_codes(first.get("JASDM")),
            historical_rooms=_codes(first.get("LSJASDM")), prefer=_text(first.get("Prefer_Time")),
            unavailable=_text(first.get("unavailable_Time")), special=_text(first.get("PKYQMS")),
            max_block=int(max_block), block_template=template,
            check_room=_check_flag(first.get("IF_ROOM_CONFICT"), "IF_ROOM_CONFICT", issues),
            check_class=check_class,
            issues=issues, total_hours=total_hours,
        )
        dataset.courses[identifier] = course


def _load_rooms(rows: list[tuple[int, dict[str, Any]]], dataset: Dataset) -> None:
    for number, row in rows:
        identifier = _text(row.get("JASDM"))
        if not identifier:
            dataset.issues.append(f"metadata:room row {number}: missing JASDM")
            continue
        problems: list[str] = []
        capacity = _number(row.get("SKZWS"), "SKZWS", problems)
        campus = _text(row.get("MC"))
        if not campus:
            problems.append("metadata: missing campus")
        flag = _text(row.get("SFYXPK"))
        if flag not in ("0", "0.0", "1", "1.0"):
            problems.append(f"metadata: invalid SFYXPK: {flag!r}")
        room = Room(id=identifier, name=_text(row.get("JASMC")), campus=campus,
                    capacity=capacity, type=_text(row.get("JASLX")),
                    building=_text(row.get("JXLDM")), enabled=flag in ("1", "1.0") and not problems)
        if identifier in dataset.rooms and dataset.rooms[identifier] != room:
            dataset.rooms[identifier].enabled = False
            problems.append("metadata: contradictory duplicate room; disabled for new placements")
        elif identifier not in dataset.rooms:
            dataset.rooms[identifier] = room
        dataset.issues.extend(f"metadata:room {identifier!r} row {number}: {problem}" for problem in problems)


def _load_placements(rows: list[tuple[int, dict[str, Any]]], dataset: Dataset) -> None:
    seen: set[Placement] = set()
    for number, row in rows:
        identifier = _text(row.get("教学班ID"))
        room_id = _text(row.get("教室代码"))
        raw_room_id = room_id
        period_text = _normalize_time(_text(row.get("节次")))
        range_match = re.fullmatch(r"第?(\d+)-(\d+)节?", period_text)
        # A reversed period interval alone is an error. It is an empty export
        # tail only when this exact schedule row explicitly declares zero hours.
        try:
            row_hours = float(_text(row.get("周学时")))
        except ValueError:
            row_hours = None
        if row_hours == 0 and range_match and int(range_match[2]) < int(range_match[1]):
            dataset.issues.append(
                f"baseline:zero_hour_placeholder row {number} course {identifier!r}: "
                f"zero-hour reversed-period export row, no occupancy; "
                f"raw weeks={_text(row.get('上课周次'))!r}, day={_text(row.get('星期'))!r}, "
                f"periods={_text(row.get('节次'))!r}, room={raw_room_id!r}, weekly_hours=0"
            )
            continue
        problems: list[str] = []
        course = dataset.courses.get(identifier)
        if course is None:
            problems.append(f"missing course metadata {identifier!r}")
        shared_room = course is not None and not course.check_room and _normalize_time(room_id) == "(共占)"
        if shared_room:
            # Empty id explicitly means no exclusive classroom resource. Never
            # invent a room, and never apply this exception to normal courses.
            room_id = ""
        elif room_id not in dataset.rooms:
            problems.append(f"missing room metadata {room_id!r}")
        try:
            weeks = _week_list(row.get("上课周次"))
        except ValueError as error:
            weeks = frozenset()
            problems.append(str(error))
        try:
            periods = _periods(row.get("节次"))
        except ValueError as error:
            periods = ()
            problems.append(str(error))
        day_text = _text(row.get("星期"))
        day = DAYS.get(day_text, -1)
        if day < 0:
            problems.append(f"unsupported weekday: {day_text!r}")
        placement = Placement(identifier, weeks, day, periods, room_id)
        if placement not in seen:
            dataset.placements.append(placement)
            seen.add(placement)
        # Each invalid row remains in diagnostics even when its placement is a duplicate.
        if problems:
            dataset.issues.append(f"baseline:row {number} course {identifier!r}: {'; '.join(problems)}; "
                                  f"raw weeks={_text(row.get('上课周次'))!r}, day={day_text!r}, "
                                  f"periods={_text(row.get('节次'))!r}, room={raw_room_id!r}, "
                                  f"weekly_hours={_text(row.get('周学时'))!r}")


def load_dataset(course_path: str | Path, room_path: str | Path,
                 schedule_path: str | Path, class_path: str | Path | None = None) -> Dataset:
    """Read Excel sources, optionally validating classes against a BJMC registry.

    With no explicit class path, discover ``班级表.xlsx`` beside the course file.
    If absent, retain conservative token checks and report the coverage limit.
    """
    dataset = Dataset()
    paths = {"courses": Path(course_path), "rooms": Path(room_path), "schedule": Path(schedule_path)}
    inferred_class_path = paths["courses"].parent / "班级表.xlsx"
    if class_path is not None:
        paths["classes"] = Path(class_path)
    elif inferred_class_path.is_file():
        paths["classes"] = inferred_class_path
    for key, path in paths.items():
        with path.open("rb") as stream:
            digest = sha256()
            for block in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(block)
        dataset.source_hashes[key] = digest.hexdigest()
    registry = None
    if "classes" in paths:
        class_rows = _read_table(paths["classes"], False)
        if any("BJMC" not in row for _, row in class_rows):
            raise ValueError(f"class registry requires a single header row containing BJMC: {paths['classes']}")
        registry = frozenset(_text(row.get("BJMC")) for _, row in class_rows if _text(row.get("BJMC")))
    dataset.class_registry = registry
    _load_courses(_read_table(paths["courses"], True), dataset, registry)
    if registry is None and any(course.check_class for course in dataset.courses.values()):
        dataset.issues.append("coverage: class registry unavailable; class identifiers use token checks only; "
                              "unregistered class names cannot be reliably detected")
    _load_rooms(_read_table(paths["rooms"], True), dataset)
    _load_placements(_read_table(paths["schedule"], False), dataset)
    return dataset
