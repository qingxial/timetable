"""JSON tool boundary for the existing bounded timetable repair engine.

This module proposes and diagnoses; it does not publish a timetable, infer
permission from natural language, or execute generated code. Exact conflict
certificates are explicitly scoped to their finite candidate model.
"""
from __future__ import annotations

import argparse
from collections import Counter
from copy import deepcopy
from dataclasses import asdict
import hashlib
import json
import math
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

try:
    from .repair_courses import RepairConfig, preference_cost, repair
    from .repair_data import load_dataset, parse_time_windows
except ImportError:
    from repair_courses import RepairConfig, preference_cost, repair
    from repair_data import load_dataset, parse_time_windows


def _object(properties, required):
    return {"type": "object", "properties": properties, "required": required, "additionalProperties": False}


_PATH = {"type": "string", "minLength": 1, "pattern": r"\S"}
_IDS = {"type": "array", "items": {"type": "string", "minLength": 1}, "uniqueItems": True}
CONFIG_SCHEMA = _object({
    "allowed_changes": {"type": "array", "items": {"type": "string", "enum": [
        "time_preference", "building_preference", "historical_room"]}, "uniqueItems": True},
    "time_scope": {"type": "string", "enum": ["strict", "same_day", "weekdays", "configured_days"]},
    "single_period_starts": {"type": "string", "enum": ["odd", "any"]},
    "target_ids": _IDS, "locked_ids": _IDS, "movable_ids": _IDS, "reviewed_soft_time_ids": _IDS,
    "max_moved_courses": {"type": "integer", "minimum": 0, "maximum": 3},
    "max_candidates_per_course": {"type": "integer", "minimum": 1},
    "max_search_nodes": {"type": "integer", "minimum": 1},
    "max_weekly_hours": {"type": "integer", "minimum": 1, "maximum": 77},
    "max_blocks_per_day": {"type": "integer", "minimum": 1, "maximum": 11},
    "filter_fixed_occupancy": {"type": "boolean"},
    "time_limit_seconds": {"type": "number", "exclusiveMinimum": 0},
    "days": {"type": "array", "items": {"type": "integer", "minimum": 0, "maximum": 6}, "minItems": 1, "uniqueItems": True},
    "periods": {"type": "array", "items": {"type": "integer", "minimum": 1, "maximum": 11}, "minItems": 1, "uniqueItems": True},
    "additional_forbidden": {"type": "string"},
}, ["allowed_changes"])

TOOL_SCHEMAS = [
    {"type": "function", "name": "diagnose_remaining",
     "description": "Read an existing proposal and return status counts and per-course evidence. A failed bounded search is not an infeasibility proof.",
     "parameters": _object({"proposal_path": _PATH}, ["proposal_path"])},
    {"type": "function", "name": "propose_repair",
     "description": "Dry-run the existing bounded repair engine on explicit Excel sources and an explicit adjustment whitelist. Return a proposal without writing source files or authorizing publication.",
     "parameters": _object({"course_path": _PATH, "room_path": _PATH, "schedule_path": _PATH,
                            "failed_path": _PATH, "class_path": _PATH, "config": CONFIG_SCHEMA},
                           ["course_path", "room_path", "schedule_path", "failed_path", "config"])},
    {"type": "function", "name": "compare_proposals",
     "description": "Compare proposals only when baseline source hashes and target sets match; disclose policy and search-budget differences rather than treating changed constraints as an algorithm comparison.",
     "parameters": _object({"proposal_paths": {"type": "array", "items": _PATH, "minItems": 2, "uniqueItems": True}},
                           ["proposal_paths"])},
]
_NONNEGATIVE = {"type": "number", "minimum": 0}
_CORRECTION_BASE = {"course_id": {"type": "string", "minLength": 1}, "evidence": _PATH}
def _correction_schema(operation, properties, required):
    return _object({**_CORRECTION_BASE, "operation": {"type": "string", "enum": [operation]}, **properties},
                   ["course_id", "operation", "evidence", *required])
CORRECTIONS_SCHEMA = {"type": "array", "items": {"oneOf": [
    _correction_schema("redistribute_hours", {
        "expected_weekly_hours": _NONNEGATIVE, "expected_total_hours": _NONNEGATIVE,
        "authoritative_field": {"type": "string", "enum": ["LLXS"]},
        "strategy": {"type": "string", "enum": ["balanced_frontload"]},
        "allocation_unit": {"type": "integer", "enum": [1, 2]},
        "weekly_loads": {"type": "object", "additionalProperties": {"type": "integer", "minimum": 1, "maximum": 77}},
    }, ["expected_weekly_hours", "expected_total_hours", "authoritative_field"]),
    _correction_schema("set_capacity", {"expected_capacity": _NONNEGATIVE,
        "capacity": {"type": "number", "exclusiveMinimum": 0}, "enrollment_frozen": {"type": "boolean", "enum": [True]}},
        ["expected_capacity", "capacity", "enrollment_frozen"]),
    _correction_schema("replace_classes", {"expected_classes": _IDS, "classes": {**_IDS, "minItems": 1}},
        ["expected_classes", "classes"]),
    _correction_schema("set_one_off_week", {
        "expected_weekly_hours": {"type": "number", "enum": [0]},
        "expected_total_hours": {"type": "integer", "minimum": 1, "maximum": 11},
        "teaching_week": {"type": "integer", "minimum": 1, "maximum": 53},
        "weeks_confirmed": {"type": "boolean", "enum": [True]},
    }, ["expected_weekly_hours", "expected_total_hours", "teaching_week", "weeks_confirmed"]),
    _correction_schema("quarantine_incomplete_segments", {"expected_incomplete_count": {"type": "integer", "minimum": 1}},
        ["expected_incomplete_count"]),
]}}
_SOURCE_PROPERTIES = {"course_path": _PATH, "room_path": _PATH, "schedule_path": _PATH,
                      "failed_path": _PATH, "class_path": _PATH, "corrections": CORRECTIONS_SCHEMA}
_SOURCE_REQUIRED = ["course_path", "room_path", "schedule_path", "failed_path"]
OPTIMIZER_SCHEMA = _object({
    "quality_weights": _object({key: _NONNEGATIVE for key in ("time_preference", "evening", "weekend")},
                               ["time_preference", "evening", "weekend"]),
    "time_limit_seconds": {"type": "number", "exclusiveMinimum": 0},
    "max_search_nodes": {"type": "integer", "minimum": 1},
    "candidate_limit_per_course": {"type": "integer", "minimum": 1},
}, ["quality_weights"])
EXPLAIN_SCHEMA = _object({
    "max_candidates_per_course": {"type": "integer", "minimum": 1},
    "max_nodes": {"type": "integer", "minimum": 1},
    "max_oracle_calls": {"type": "integer", "minimum": 1},
    "time_limit_seconds": {"type": "number", "exclusiveMinimum": 0},
    "allow_scoped_proof": {"type": "boolean"},
}, [])
TOOL_SCHEMAS[1]["parameters"]["properties"]["corrections"] = CORRECTIONS_SCHEMA
for _name, _description, _extras, _required in [
    ("repair_with_fallbacks", "Try daytime, then evening, then weekend only within the explicit caller domain, preserving earlier successful placements; optional audited input corrections.",
     {"config": CONFIG_SCHEMA, "stages": {"type": "array", "items": {"type": "string", "enum": ["daytime", "evening", "weekend"]}, "uniqueItems": True, "minItems": 1},
      "total_time_limit_seconds": {"type": "number", "exclusiveMinimum": 0}}, ["config"]),
    ("preview_data_corrections", "Validate explicit correction operations against original values and evidence, return audit and effective data hash without writing workbooks.", {}, ["corrections"]),
    ("optimize_local_repair", "Phillips-inspired joint finite-domain search: maximize complete insertions, then minimize original-course moves and quality cost. Proofs are scoped to retained candidates.",
     {"config": CONFIG_SCHEMA, "optimizer_config": OPTIMIZER_SCHEMA}, ["config", "optimizer_config"]),
    ("quality_frontier", "Lindahl-inspired movement-budget comparison using the same candidate domain, success level and shared runtime budget; explicit evening/weekend weights.",
     {"config": CONFIG_SCHEMA, "optimizer_config": OPTIMIZER_SCHEMA,
      "movement_budgets": {"type": "array", "items": {"type": "integer", "minimum": 0, "maximum": 3}, "minItems": 1, "uniqueItems": True}}, ["config", "optimizer_config", "movement_budgets"]),
    ("explain_conflict", "QuickXplain over an exact finite-candidate oracle with SAT/UNSAT/UNKNOWN, background checks and deletion verification. Releasing occupancy is an explanation, not a complete movement plan.",
     {"config": CONFIG_SCHEMA, "explain_config": EXPLAIN_SCHEMA}, ["config", "explain_config"]),
    ("plan_course_corrections", "Return specific per-course correction actions, source rows, proposed values, required evidence and revalidation conditions, including optional skipped-course records.",
     {"proposal_path": _PATH, "source_workbook": _PATH, "skipped_path": _PATH}, ["proposal_path"]),
    ("extend_repair_proposal", "Validate and preserve every successful insertion in an existing no-move proposal, then try an explicit additional target list using day/evening/weekend stages. The seed corrections must be an exact prefix of the supplied corrections.",
     {"seed_proposal_path": _PATH, "additional_target_ids": {**_IDS, "minItems": 1},
      "config": CONFIG_SCHEMA,
      "stages": {"type": "array", "items": {"type": "string", "enum": ["daytime", "evening", "weekend"]}, "uniqueItems": True, "minItems": 1},
      "total_time_limit_seconds": {"type": "number", "exclusiveMinimum": 0}},
     ["seed_proposal_path", "additional_target_ids", "config"]),
]:
    TOOL_SCHEMAS.append({"type": "function", "name": _name, "description": _description,
                         "parameters": _object({**_SOURCE_PROPERTIES, **_extras}, [*_SOURCE_REQUIRED, *_required])})


REQUEST_SCHEMA = _object({"name": {"type": "string", "minLength": 1}, "arguments": {"type": "object"}},
                         ["name", "arguments"])


class ToolError(ValueError):
    def __init__(self, code: str, message: str, details: Any = None):
        super().__init__(message)
        self.code, self.details = code, details


def _validate(value, schema, location="arguments"):
    """Validate the small, explicit JSON Schema vocabulary published above."""
    if "oneOf" in schema:
        matches = 0
        for candidate in schema["oneOf"]:
            try:
                _validate(value, candidate, location)
                matches += 1
            except ToolError:
                pass
        if matches != 1:
            raise ToolError("invalid_arguments", f"{location} must match exactly one supported operation schema")
        return
    kind = schema.get("type")
    try:
        finite_number = type(value) in (int, float) and math.isfinite(value)
    except OverflowError:
        finite_number = False
    valid = {
        "object": isinstance(value, dict), "array": isinstance(value, list),
        "string": isinstance(value, str), "integer": type(value) is int,
        "number": finite_number,
        "boolean": type(value) is bool,
    }.get(kind, False)
    if not valid:
        raise ToolError("invalid_arguments", f"{location} must be {kind}")
    if "enum" in schema and value not in schema["enum"]:
        raise ToolError("invalid_arguments", f"{location} must be one of {schema['enum']!r}")
    if kind == "object":
        missing = set(schema.get("required", ())) - value.keys()
        unknown = value.keys() - schema.get("properties", {}).keys()
        if missing:
            raise ToolError("invalid_arguments", f"{location} missing required fields: {sorted(missing)}")
        if schema.get("additionalProperties") is False and unknown:
            raise ToolError("invalid_arguments", f"{location} unknown fields: {sorted(unknown)}")
        for key, item in value.items():
            if key in schema.get("properties", {}):
                _validate(item, schema["properties"][key], f"{location}.{key}")
            elif isinstance(schema.get("additionalProperties"), dict):
                _validate(item, schema["additionalProperties"], f"{location}.{key}")
    elif kind == "array":
        if len(value) < schema.get("minItems", 0):
            raise ToolError("invalid_arguments", f"{location} has too few items")
        if schema.get("uniqueItems") and len({json.dumps(v, sort_keys=True) for v in value}) != len(value):
            raise ToolError("invalid_arguments", f"{location} contains duplicate items")
        for index, item in enumerate(value):
            _validate(item, schema["items"], f"{location}[{index}]")
    elif kind == "string":
        if len(value) < schema.get("minLength", 0) or ("pattern" in schema and not re.search(schema["pattern"], value)):
            raise ToolError("invalid_arguments", f"{location} is empty or malformed")
    elif kind in ("integer", "number"):
        if ("minimum" in schema and value < schema["minimum"] or
            "maximum" in schema and value > schema["maximum"] or
            "exclusiveMinimum" in schema and value <= schema["exclusiveMinimum"]):
            raise ToolError("invalid_arguments", f"{location} is outside the allowed range")


def _read_json(path):
    def duplicate_keys(pairs):
        result = {}
        for key, value in pairs:
            if key in result:
                raise ToolError("input_error", f"Duplicate JSON field: {key}")
            result[key] = value
        return result

    def invalid_constant(value):
        raise ToolError("input_error", f"Non-finite JSON number: {value}")

    def finite_float(value):
        result = float(value)
        if not math.isfinite(result):
            invalid_constant(value)
        return result

    with Path(path).open(encoding="utf-8") as stream:
        return json.load(stream, object_pairs_hook=duplicate_keys, parse_constant=invalid_constant,
                         parse_float=finite_float)


def _hash(path):
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _read_failed(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = list(next(rows, ()))
        key = "jxbid" if "jxbid" in headers else "教学班ID"
        if key not in headers:
            raise ToolError("input_error", "Failed-course workbook requires jxbid or 教学班ID header")
        index = headers.index(key)
        values = []
        for number, row in enumerate(rows, start=2):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            raw = row[index] if index < len(row) else None
            if raw is None or not str(raw).strip():
                raise ToolError("input_error", f"Failed-course row {number} has no course ID")
            values.append(str(raw).strip())
        return list(dict.fromkeys(values))
    finally:
        workbook.close()


def _proposal(path):
    proposal = _read_json(path)
    # CLI responses can be consumed directly, as can legacy proposal.json files.
    if isinstance(proposal, dict) and proposal.get("ok") is True and proposal.get("tool") in {"propose_repair", "repair_with_fallbacks", "optimize_local_repair", "extend_repair_proposal"}:
        envelope_result = proposal.get("result")
        proposal = envelope_result.get("proposal") if isinstance(envelope_result, dict) else None
    if (not isinstance(proposal, dict) or proposal.get("mode") != "proposal"
            or type(proposal.get("schema_version")) is not int or proposal["schema_version"] != 1):
        raise ToolError("invalid_proposal", f"Unsupported proposal format: {path}")
    for key, kind in (("source_hashes", dict), ("config", dict), ("summary", dict), ("results", list), ("changes", list)):
        if not isinstance(proposal.get(key), kind):
            raise ToolError("invalid_proposal", f"Proposal {key} has the wrong type: {path}")
    for key in ("courses", "rooms", "schedule"):
        if not re.fullmatch(r"[0-9a-f]{64}", str(proposal["source_hashes"].get(key, ""))):
            raise ToolError("invalid_proposal", f"Proposal lacks a valid {key} source SHA-256: {path}")
    if any(not isinstance(value, str) or not re.fullmatch(r"[0-9a-f]{64}", value)
           for value in proposal["source_hashes"].values()):
        raise ToolError("invalid_proposal", f"Malformed source SHA-256: {path}")
    for key in ("requested", "inserted", "moved"):
        if type(proposal["summary"].get(key)) is not int or proposal["summary"][key] < 0:
            raise ToolError("invalid_proposal", f"Proposal summary.{key} must be a non-negative integer")
    config = deepcopy(proposal["config"])
    # Existing repair output always has this whitelist; do not infer it from text.
    _validate(config, CONFIG_SCHEMA, "proposal.config")
    targets = []
    for result in proposal["results"]:
        if (not isinstance(result, dict) or not isinstance(result.get("course_id"), str)
                or not result["course_id"] or not isinstance(result.get("status"), str)
                or not isinstance(result.get("diagnosis", {}), dict)):
            raise ToolError("invalid_proposal", f"Malformed per-course result: {path}")
        targets.append(result["course_id"])
    if len(set(targets)) != len(targets):
        raise ToolError("invalid_proposal", f"Duplicate result course IDs: {path}")
    if proposal["summary"].get("requested") != len(targets):
        raise ToolError("invalid_proposal", f"Requested count differs from result target set: {path}")
    return proposal


def _status_evidence(status):
    if status == "placed":
        return "feasible_incremental_proposal"
    if status == "already_scheduled":
        return "baseline_record_exists_not_full_validation"
    if status in ("search_limit", "not_found_within_limits", "unknown"):
        return "unknown_within_search_limits"
    if status.lower() in ("infeasible", "unsat"):
        return "unverified_infeasibility_claim"
    if status in ("no_matching_room", "no_time_pattern"):
        return "no_candidate_in_current_model"
    return "unresolved_data_policy_or_model_condition"


def _diagnose(arguments):
    proposal = _proposal(arguments["proposal_path"])
    rows = [{**deepcopy(row), "evidence_status": _status_evidence(row["status"]),
             "infeasibility_proven": False} for row in proposal["results"]]
    return {"source_hashes": proposal["source_hashes"], "policy": proposal["config"],
            "summary": proposal["summary"], "status_counts": dict(Counter(r["status"] for r in rows)),
            "remaining_count": sum(r["status"] not in ("placed", "already_scheduled") for r in rows),
            "courses": rows, "validation": proposal.get("validation", {}),
            "data_issues": proposal.get("data_issues", []),
            "interpretation": "Statuses are evidence from a bounded heuristic and its input model; no UNSAT certificate or global infeasibility proof is provided."}


def _propose(arguments):
    return _run_repair_tool("propose_repair", arguments)


def _run_repair_tool(name, arguments):
    try:
        from .repair_workflows import apply_data_corrections, repair_with_fallbacks
        from .repair_optimizer import optimize_local, quality_frontier
        from .repair_explanations import explain_conflict
        from .repair_action_plans import build_action_plans, load_action_plan_supplemental
        from .repair_extensions import extend_repair_proposal
    except ImportError:
        from repair_workflows import apply_data_corrections, repair_with_fallbacks
        from repair_optimizer import optimize_local, quality_frontier
        from repair_explanations import explain_conflict
        from repair_action_plans import build_action_plans, load_action_plan_supplemental
        from repair_extensions import extend_repair_proposal
    paths = {key: Path(arguments[f"{key[:-1] if key in ('courses', 'rooms') else key}_path"])
             for key in ("courses", "rooms", "schedule", "failed")}
    class_path = Path(arguments["class_path"]) if "class_path" in arguments else paths["courses"].parent / "班级表.xlsx"
    if "class_path" in arguments or class_path.is_file():
        paths["classes"] = class_path
    for key in ("source_workbook", "skipped_path"):
        if key in arguments:
            paths[key] = Path(arguments[key])
    before = {key: _hash(path) for key, path in paths.items()}
    data = load_dataset(paths["courses"], paths["rooms"], paths["schedule"], paths.get("classes"))
    targets = _read_failed(paths["failed"])
    data.source_hashes["failed"] = before["failed"]
    seed = None
    if name == "extend_repair_proposal":
        seed = _proposal(arguments["seed_proposal_path"])
        _metrics(seed)
        seed_corrections = seed.get("correction_parameters", [])
        supplied = arguments.get("corrections", [])
        if supplied[:len(seed_corrections)] != seed_corrections:
            raise ToolError("invalid_arguments", "Seed correction parameters must be an exact prefix of the supplied corrections")
        extra_ids = set(arguments["additional_target_ids"])
        if any(item['course_id'] not in extra_ids for item in supplied[len(seed_corrections):]):
            raise ToolError("invalid_arguments", "Additional corrections may only affect additional targets")
        seed_data, _ = apply_data_corrections(data, seed_corrections)
        if seed_data.source_hashes != seed["source_hashes"]:
            raise ToolError("incomparable_proposals", "Seed source/correction hashes do not match the replayed dataset")
    data, audit = apply_data_corrections(data, arguments.get("corrections", []))
    config = deepcopy(arguments.get("config", {"allowed_changes": []}))
    policy = RepairConfig(**config); policy.validate()
    outside = set(policy.target_ids) - set(arguments["additional_target_ids"] if name == "extend_repair_proposal" else targets)
    if outside:
        raise ToolError("invalid_arguments", "config.target_ids contains IDs outside the supplied failed-course workbook", sorted(outside))
    if name == "preview_data_corrections":
        output = {"audit": audit, "effective_source_hashes": data.source_hashes,
                  "mode": "correction_preview", "source_files_unchanged": True}
    elif name == "quality_frontier":
        output = quality_frontier(data, targets, config, arguments["optimizer_config"], arguments["movement_budgets"])
    elif name == "explain_conflict":
        output = explain_conflict(data, policy.target_ids or targets, config, arguments["explain_config"])
    elif name == "plan_course_corrections":
        proposal = _proposal(arguments["proposal_path"])
        for key in ("courses", "rooms", "schedule", "classes", "corrections"):
            if proposal["source_hashes"].get(key) != data.source_hashes.get(key):
                raise ToolError("incomparable_proposals", "Action plan input differs from the proposal's effective dataset", key)
        supplemental = load_action_plan_supplemental(arguments.get("source_workbook"), paths["courses"], arguments.get("skipped_path"), paths["schedule"])
        output = build_action_plans(data, proposal, supplemental)
    else:
        if name == "extend_repair_proposal":
            proposal = extend_repair_proposal(data, seed, arguments["additional_target_ids"], config,
                arguments.get("stages"), arguments.get("total_time_limit_seconds", 180))
        elif name == "repair_with_fallbacks":
            proposal = repair_with_fallbacks(data, targets, config, arguments.get("stages"), arguments.get("total_time_limit_seconds", 180))
        elif name == "optimize_local_repair":
            proposal = optimize_local(data, targets, config, arguments["optimizer_config"])
        else:
            proposal = repair(data, targets, config)
        proposal["input_corrections"] = audit
        proposal["correction_parameters"] = deepcopy(arguments.get("corrections", []))
        output = {"proposal": proposal, "policy": proposal["config"],
                  "source_files_unchanged": True, "mode": "dry_run_proposal",
                  "authorization_basis": "Only explicit caller parameters are executed; audit records retain original values and planning assumptions."}
    after = {key: _hash(path) for key, path in paths.items()}
    if before != after or any(data.source_hashes.get(key) != before[key] for key in ("courses", "rooms", "schedule", "failed")):
        raise ToolError("input_changed", "Source files changed during the run; result rejected")
    output["source_files_unchanged"] = True
    output["input_corrections"] = audit
    return output


def _metrics(proposal):
    operations = {"insert": set(), "move": set()}
    cost = 0
    seen = set()
    for change in proposal["changes"]:
        if (not isinstance(change, dict) or not isinstance(change.get("course_id"), str)
                or change.get("operation") not in operations or not isinstance(change.get("after"), list)
                or not change["after"]
                or not isinstance(change.get("original_time_preference", ""), str)):
            raise ToolError("invalid_proposal", "Malformed change record")
        cid = change["course_id"]
        if cid in seen:
            raise ToolError("invalid_proposal", f"Duplicate change record for {cid}")
        seen.add(cid)
        operations[change["operation"]].add(cid)
        pref = parse_time_windows(change.get("original_time_preference", ""))
        patterns = set()
        for placement in change["after"]:
            if (not isinstance(placement, dict) or type(placement.get("day")) is not int
                    or placement["day"] not in range(7) or not isinstance(placement.get("periods"), list)
                    or not placement["periods"] or any(type(p) is not int or p not in range(1, 12) for p in placement["periods"])):
                raise ToolError("invalid_proposal", f"Malformed changed placement for {cid}")
            if placement["periods"] != sorted(set(placement["periods"])):
                raise ToolError("invalid_proposal", f"Changed periods must be sorted and unique for {cid}")
            patterns.add((placement["day"], tuple(placement["periods"])))
        cost += sum(preference_cost(day, periods, pref) for day, periods in patterns)
    inserted, moved = len(operations["insert"]), len(operations["move"])
    if proposal["summary"].get("inserted") != inserted or proposal["summary"].get("moved") != moved:
        raise ToolError("invalid_proposal", "Summary insert/move counts disagree with unique change records")
    if {r["course_id"] for r in proposal["results"] if r["status"] == "placed"} != operations["insert"]:
        raise ToolError("invalid_proposal", "Placed target results disagree with inserted change records")
    return {"inserted": inserted, "moved_original_courses": moved, "preference_cost": cost,
            "inserted_ids": sorted(operations["insert"]), "moved_ids": sorted(operations["move"])}


def _compare(arguments):
    proposals = [_proposal(path) for path in arguments["proposal_paths"]]
    # A failed-list path may differ between exports. The target-set equality
    # check below protects that dimension; baseline resource content must match.
    def baseline(proposal):
        return {key: proposal["source_hashes"].get(key) for key in ("courses", "rooms", "schedule", "classes", "corrections")}

    reference = baseline(proposals[0])
    targets = {r["course_id"] for r in proposals[0]["results"]}
    for proposal in proposals[1:]:
        if baseline(proposal) != reference:
            raise ToolError("incomparable_proposals", "Baseline course/room/schedule/class/correction hashes differ")
        if {r["course_id"] for r in proposal["results"]} != targets:
            raise ToolError("incomparable_proposals", "Proposal target course sets differ")
    configs = [asdict(RepairConfig(**p["config"])) for p in proposals]
    differing = {key: [config[key] for config in configs] for key in configs[0]
                 if any(config[key] != configs[0][key] for config in configs[1:])}
    rows = [{"path": path, **_metrics(proposal), "policy": config,
             "status_counts": dict(Counter(r["status"] for r in proposal["results"])),
             "validation": proposal.get("validation", {})}
            for path, proposal, config in zip(arguments["proposal_paths"], proposals, configs)]
    return {"baseline_source_hashes": reference, "target_ids": sorted(targets), "proposals": rows,
            "policy_or_budget_differences": differing, "same_recorded_policy_and_budget": not differing,
            "preference_cost_definition": "Sum of the existing engine's preference_cost for each unique (day, periods) pattern after each changed course; includes inserted and moved courses, without multiplying by weeks.",
            "interpretation": "Compare these as conditional policy/search scenarios. Changed policies do not establish better performance under identical hard constraints; identical recorded config alone is not a controlled solver experiment."}


def _error(name, code, message, details=None):
    result = {"ok": False, "tool": name, "error": {"code": code, "message": message}}
    if details is not None:
        result["error"]["details"] = details
    return result


def dispatch_tool(name, arguments):
    """Return one JSON-compatible success/error envelope; never write a file."""
    schemas = {schema["name"]: schema["parameters"] for schema in TOOL_SCHEMAS}
    if not isinstance(name, str) or name not in schemas:
        return _error(name if isinstance(name, str) else None, "unknown_tool", "Unknown timetable tool")
    try:
        _validate(arguments, schemas[name])
        handlers = {"diagnose_remaining": _diagnose, "propose_repair": _propose, "compare_proposals": _compare}
        for tool_name in schemas.keys() - handlers.keys():
            handlers[tool_name] = lambda args, selected=tool_name: _run_repair_tool(selected, args)
        return {"ok": True, "tool": name, "result": handlers[name](deepcopy(arguments))}
    except ToolError as error:
        return _error(name, error.code, str(error), error.details)
    except (OSError, ValueError, TypeError, KeyError) as error:
        return _error(name, "input_error", str(error))
    except RuntimeError as error:
        return _error(name, "repair_rejected", str(error))
    except Exception as error:
        # Keep the tool protocol intact for malformed workbooks (e.g. corrupt
        # ZIP/XML) and unexpected implementation errors. Never continue repair.
        return _error(name, "tool_error", f"{type(error).__name__}: {error}")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--request", required=True, type=Path, help='JSON object with exactly "name" and "arguments"')
    parser.add_argument("--output", required=True, type=Path, help="New JSON output file; existing files are never overwritten")
    arguments = parser.parse_args(argv)
    if arguments.output.exists():
        print(json.dumps(_error(None, "output_exists", "Output already exists; choose a new file"), ensure_ascii=False))
        return 1
    try:
        request = _read_json(arguments.request)
        _validate(request, REQUEST_SCHEMA, "request")
        response = dispatch_tool(request["name"], request["arguments"])
    except (ToolError, OSError, ValueError) as error:
        response = _error(None, getattr(error, "code", "input_error"), str(error))
    try:
        with arguments.output.open("x", encoding="utf-8") as stream:
            json.dump(response, stream, ensure_ascii=False, indent=2, allow_nan=False)
            stream.write("\n")
    except OSError as error:
        print(json.dumps(_error(response.get("tool"), "output_error", str(error)), ensure_ascii=False))
        return 1
    print(json.dumps({"ok": response["ok"], "tool": response["tool"], "output": str(arguments.output)}, ensure_ascii=False))
    return 0 if response["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
