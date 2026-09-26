"""Evidence-backed, per-course repair plans; never modify data or infer consent.

``build_action_plans(dataset, proposal, supplemental=None)`` is a read-only JSON
boundary. Optional workbook enrichment records exact source rows/cells without
including teacher names, teacher identifiers, or student rosters.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import replace
import hashlib
import math
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

try:
    from .repair_courses import Occupancy, RepairConfig, effective_changes, room_rejections, time_options
    from .repair_data import Placement
except ImportError:
    from repair_courses import Occupancy, RepairConfig, effective_changes, room_rejections, time_options
    from repair_data import Placement


ORIGINAL_FIELDS = ("课程序号", "课程名称", "课程类别名称", "课程属性", "校区", "教学班名称",
                   "行政班", "人数上限", "实际人数", "周数", "周课时", "起始周", "结束周", "教室类型", "理论学时")
CONVERTED_FIELDS = ("JXBID", "KCM", "ZXS", "LLXS", "SKXS", "XS", "SKZCDM", "RWJSZCDM",
                    "KRL", "SKXQ", "JASLXMC", "JXLDM", "JASDM", "LSJASDM", "TJBJ",
                    "SFXYPK", "SFXYJAS", "IF_ROOM_CONFICT", "IF_CLASS_CONFICT")
SCHEDULE_FIELDS = ("教学班ID", "课程名称", "周学时", "上课周次", "星期", "节次", "教室代码")
SKIPPED_FIELDS = ("教学班ID", "课程名称", "跳过主因", "SFXYPK", "SFXYJAS", "ZXS", "LLXS", "SKZCDM")


def _text(value):
    return "" if value is None else str(value).strip()


def _number(value):
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def _read_records(path, fields, id_field, sheet=None, header_row=1):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        rows = worksheet.iter_rows(values_only=True)
        for _ in range(header_row - 1):
            next(rows)
        headers = [_text(v) for v in next(rows)]
        if id_field not in headers:
            raise ValueError(f"Missing {id_field} header in {path}")
        result = defaultdict(list)
        columns = {name: headers.index(name) for name in fields if name in headers}
        for number, row in enumerate(rows, header_row + 1):
            values = {name: _text(row[index]) if index < len(row) else "" for name, index in columns.items()}
            identifier = values.get(id_field, "")
            if not identifier:
                continue
            result[identifier].append({"file": str(Path(path)), "sheet": worksheet.title, "row": number,
                "values": values, "cells": {name: f"{get_column_letter(index+1)}{number}" for name, index in columns.items()}})
        return dict(result)
    finally:
        workbook.close()


def load_action_plan_supplemental(source_path=None, converted_path=None, skipped_path=None, schedule_path=None):
    """Read caller-supplied workbooks. No default machine-specific paths or writes."""
    result = {"original_courses": {}, "converted_courses": {}, "skipped_courses": {}, "schedule_rows": {},
              "source_hashes": {}}
    specs = (("original_courses", source_path, ORIGINAL_FIELDS, "课程序号", "教学任务", 1),
             ("converted_courses", converted_path, CONVERTED_FIELDS, "JXBID", None, 2),
             ("skipped_courses", skipped_path, SKIPPED_FIELDS, "教学班ID", "逐课明细", 1),
             ("schedule_rows", schedule_path, SCHEDULE_FIELDS, "教学班ID", None, 1))
    for key, path, fields, identifier, sheet, header in specs:
        if path is None:
            continue
        path = Path(path)
        before = hashlib.sha256(path.read_bytes()).hexdigest()
        result[key] = _read_records(path, fields, identifier, sheet, header)
        if before != hashlib.sha256(path.read_bytes()).hexdigest():
            raise ValueError(f"Supplemental source changed during read: {path}")
        result["source_hashes"][key] = before
    return result


def _action(code, fields, values, evidence, changes, required, revalidate, status="needs_evidence", auto=False):
    return {"code": code, "problem_fields": fields, "current_values": values, "evidence": evidence,
            "proposed_changes": changes, "required_data_or_confirmation": required,
            "execution_status": status, "auto_executable": auto, "revalidation_conditions": revalidate}


def _segment_candidate(total, weeks):
    """Arithmetic allocation only. Never invent which calendar weeks get a load."""
    if not weeks or total is None or total <= 0 or abs(total-round(total)) > 1e-5:
        return None
    count, total = len(weeks), int(round(total))
    average = total / count
    # Prefer normal two-period units when every week can remain active.
    low = 2 * math.floor(average / 2)
    step = 2 if low >= 2 and total % 2 == 0 else 1
    if step == 1:
        low = math.floor(average)
    high = low + step
    high_count = (total-low*count) // step
    low_count = count-high_count
    segments = [{"weekly_hours": hours, "week_count": n, "assigned_weeks": None}
                for hours, n in ((low, low_count), (high, high_count)) if n]
    if len(segments) == 1:
        # A constant load uses the already supplied week set; no new date choice.
        segments[0]["assigned_weeks"] = sorted(weeks)
    return {"reference_total_hours": total, "existing_week_domain": sorted(weeks), "segments": segments,
            "arithmetic_total": sum(x["weekly_hours"]*x["week_count"] for x in segments),
            "calendar_assignment_complete": len(segments) == 1,
            "preserves_every_week_as_active": all(x["weekly_hours"] > 0 for x in segments),
            "requires_confirmation": True,
            "note": "Counts are a candidate, not inferred historical teaching dates. Confirm LLXS authority and assign exact weeks before execution."}


def _source_evidence(cid, supplemental):
    original = supplemental.get("original_courses", {}).get(cid.split("@", 1)[0], [])
    converted = supplemental.get("converted_courses", {}).get(cid, [])
    return {"original": original, "converted": converted,
            "original_join": "base_id" if "@" in cid else "exact_id"}


def _first_values(records):
    return records[0].get("values", {}) if records else {}


def _hours_action(course, sources, reason):
    source = _first_values(sources["original"])
    candidate = _segment_candidate(course.total_hours, course.weeks)
    evidence = {"source_records": sources,
                "original_total_hours": _number(source.get("理论学时")),
                "converted_total_hours": course.total_hours,
                "original_weekly_hours": _number(source.get("周课时")),
                "converted_weekly_hours": course.hours,
                "converted_week_count": len(course.weeks),
                "weekly_hours_times_weeks": course.hours*len(course.weeks),
                "original_reported_week_count": _number(source.get("周数")),
                "original_start_week": _number(source.get("起始周")),
                "original_end_week": _number(source.get("结束周")),
                "original_and_converted_total_agree": bool(source) and _number(source.get("理论学时")) == course.total_hours,
                "original_and_converted_weekly_agree": bool(source) and _number(source.get("周课时")) == course.hours}
    if "@" in course.id:
        candidate = None
    proposed = [{"field": "weekly_load_by_week", "candidate_value": candidate,
        "condition": "Adopt confirmed LLXS and existing teaching-week domain as the authoritative totals; do not increase LLXS to make a schedule fit."}]
    if candidate and candidate["calendar_assignment_complete"]:
        proposed.append({"field": "ZXS", "candidate_value": candidate["segments"][0]["weekly_hours"],
            "condition": "Only after adopting the confirmed LLXS/week-domain basis; the same weekly value then applies to every existing teaching week."})
    return _action("reconcile_total_and_weekly_load", ["LLXS", "ZXS", "SKZCDM"],
        {"LLXS": course.total_hours, "ZXS": course.hours, "weeks": sorted(course.weeks)}, evidence,
        proposed,
        ["确认教学大纲总学时与授课周口径；原表同时存在周学时和总学时，数值相同不证明哪一项唯一正确。",
         "逐个指定不同负荷对应的实际教学周；不能默认把高负荷分配到前几周。",
         "虚班需要确认各段的总学时分配，不能把母课程LLXS重复分给每个虚班。"] if "@" in course.id else
        ["确认以LLXS为总量且保持现有教学周的修正策略。", "多段负荷需给出每段实际周次；零负荷段还需确认相应周是否应退出教学周。"],
        ["所有分段学时之和等于已确认LLXS；教学周与教师责任周完整覆盖。",
         "按不同周负荷排入后，对同一整批候选重新校验教师、班级、教室、容量、校区、类型和禁排。"],
        status="needs_confirmation")


def _room_action(course, dataset, config, sources):
    changes = effective_changes(course, config)
    current = [r for r in dataset.rooms.values() if not room_rejections(course, r, changes)]
    if current:
        return None
    original = _first_values(sources["original"])
    actual = _number(original.get("实际人数"))
    alternative = [] if actual is None else [r.id for r in dataset.rooms.values()
        if not room_rejections(replace(course, capacity=actual), r, changes)]
    same_campus_type = [r for r in dataset.rooms.values() if r.enabled and r.campus == course.campus
                        and (not course.room_type or r.type == course.room_type)]
    candidates = []
    if actual is not None:
        candidates.append({"field": "KRL", "candidate_value": actual,
            "condition": "Only after confirming these are frozen enrollment counts and the approved room-capacity basis.",
            "static_room_candidates": sorted(alternative), "static_candidate_count": len(alternative)})
    candidates.append({"field": "room_inventory", "candidate_value": {
        "campus": course.campus, "minimum_capacity": course.capacity, "required_type": course.room_type},
        "condition": "Provide an actually available room with the required equipment/type; do not relabel an incompatible room."})
    return _action("resolve_static_room_supply", ["KRL", "SKXQ", "JASLXMC", "JASDM", "JXLDM"],
        {"KRL": course.capacity, "campus": course.campus, "required_type": course.room_type,
         "explicit_rooms": sorted(course.explicit_rooms), "buildings": sorted(course.buildings)},
        {"source_records": sources["original"], "original_limit": _number(original.get("人数上限")),
         "actual_enrollment": actual, "current_static_room_count": 0,
         "same_campus_type_room_count": len(same_campus_type),
         "same_campus_type_max_capacity": max((r.capacity for r in same_campus_type), default=None)},
        candidates, ["人数口径需要冻结人数依据；未经确认不能降低容量需求。",
                     "新增/替换场地需提供真实容量、设备类型、校区及可用时段；若分班需另有班级、教师和课时分配。"],
        ["静态教室候选非空只是必要条件，仍需在保留既有候选的占用状态下联合试排。"])


def _baseline_rows_evidence(cid, dataset, supplemental):
    rows = supplemental.get("schedule_rows", {}).get(cid, [])
    invalid = []
    for row in rows:
        values = row["values"]
        match = re.fullmatch(r"第(\d+)-(\d+)节", values.get("节次", ""))
        if not values.get("上课周次", "").strip() or not match or int(match[1]) > int(match[2]):
            invalid.append(row)
    cells = {(w, p.day, period) for p in dataset.placements if p.course_id == cid
             for w in p.weeks for period in p.periods}
    target = dataset.courses[cid].total_hours if cid in dataset.courses else None
    return {"blocking_course_id": cid, "invalid_rows": invalid,
            "valid_distinct_periods": len(cells), "declared_LLXS": target,
            "valid_rows_already_match_LLXS": target is not None and len(cells) == target,
            "source_records": _source_evidence(cid, supplemental)}


def _resource_action(course, dataset, state, config, supplemental):
    domain, _ = time_options(course, config)
    weekly = []
    uncertain = set()
    blockers = set()
    for week in sorted(course.weeks):
        available = []
        for day, period in sorted(domain):
            occupied = set()
            for teacher, responsibility in course.teachers.items():
                if week in responsibility:
                    occupied.update(state.cells.get(("teacher", teacher, week, day, period), set()) - {course.id})
            blockers.update(occupied)
            uncertain.update(occupied & state.uncertain_ids)
            if not occupied:
                available.append([day, period])
        weekly.append({"week": week, "teacher_free_period_count": len(available),
                       "teacher_free_periods": available,
                       "needed_weekly_hours": course.hours,
                       "necessary_shortfall_in_current_fixed_state": max(0, course.hours-len(available))})
    if uncertain:
        evidence = [_baseline_rows_evidence(cid, dataset, supplemental) for cid in sorted(uncertain)]
        return _action("resolve_uncertain_baseline_rows", ["baseline.上课周次", "baseline.节次"],
            {"blocked_course_id": course.id, "uncertain_blocking_course_ids": sorted(uncertain)}, evidence,
            [{"field": "baseline_rows", "candidate_value": {
                "blocking_course_id": item["blocking_course_id"],
                "review_rows": [{"file": r["file"], "sheet": r["sheet"], "row": r["row"]} for r in item["invalid_rows"]],
                "operation_options": ["delete_confirmed_zero_contribution_export_tail", "correct_weeks_and_periods_from_authoritative_schedule"]},
                "condition": "Remove only a confirmed export remainder; matching valid-row LLXS is supporting evidence, not permission to discard a genuine lesson."}
             for item in evidence],
            ["确认缺周记录是导出残余还是实际授课；提供原始分段/教务课表证据。"],
            ["清理后的有效课时仍等于已确认LLXS，所有真实授课责任周不丢失。",
             "重建占用并重新校验整个候选批次，不能直接取消教师冲突检查。"], status="needs_confirmation")
    return _action("resolve_fixed_resource_blockers", ["teacher_availability", "approved_movable_ids", "teacher_assignment"],
        {"unavailable_Time": course.unavailable, "search_days": config.days, "search_periods": config.periods},
        {"teacher_count": len(course.teachers), "weekly_teacher_supply": weekly,
         "blocking_course_ids": sorted(blockers), "scope": "Current fixed proposal and configured days/periods, not a global infeasibility proof."},
        [{"field": "approved_movable_ids", "candidate_value": sorted(blockers),
          "condition": "A blocker inventory for selecting a small authorized neighborhood, not permission to move all listed courses."},
         {"field": "confirmed_teacher_availability_or_assignment", "candidate_value": None,
          "required_values": "For weeks with a shortage, supply actual additional available periods or a qualified replacement teacher and responsibility weeks."}],
        ["提供可调整既有课程的明确范围，或教师确实可用的新时段/代课安排。不得直接删除硬禁排。"],
        ["被移动/换教师的全部课及新增课需同时排全；原周次、总学时、容量和设备约束均保持。"])


def build_action_plans(dataset, proposal, supplemental=None):
    """Return one actionable but non-executing plan per failed or skipped ID."""
    supplemental = supplemental or {}
    if proposal.get("ok") is True and proposal.get("tool") == "propose_repair":
        proposal = proposal["result"]["proposal"]
    config = RepairConfig(**proposal.get("config", {}))
    placements = [Placement(x["course_id"], frozenset(x["weeks"]), x["day"], tuple(x["periods"]), x["room_id"])
                  for x in proposal.get("placements", [])]
    state = Occupancy(replace(dataset, placements=placements or list(dataset.placements)))
    failed = {r["course_id"]: r for r in proposal.get("results", []) if r["status"] not in {"placed", "already_scheduled"}}
    skipped = dict(supplemental.get("skipped_courses", {}))
    promoted = set(proposal.get('scope_extension', {}).get('additional_target_ids', []))
    result_ids = {r['course_id'] for r in proposal.get('results', [])}
    if not promoted <= result_ids:
        raise ValueError('Promoted skipped IDs must have explicit proposal outcomes')
    for cid in promoted:
        skipped.pop(cid, None)
    if set(failed) & set(skipped):
        raise ValueError("A course cannot be both a failed target and silently skipped")
    plans = []
    for cid in sorted(set(failed) | set(skipped)):
        course = dataset.courses.get(cid)
        sources = _source_evidence(cid, supplemental)
        raw = _first_values(sources["converted"])
        skip_values = _first_values(skipped.get(cid, []))
        status = failed[cid]["status"] if cid in failed else "silently_skipped"
        actions = []
        if course is None:
            actions.append(_action("restore_course_metadata", ["JXBID"], {"JXBID": cid}, sources, [],
                ["提供该精确教学班ID的课程、教师责任周、班级及教室需求。"], ["精确ID成功连接且无缺失关键字段。"] ))
            category = "missing_metadata"
        elif cid in skipped:
            flag = _text(skip_values.get("SFXYPK", raw.get("SFXYPK", "")))
            if flag == "0":
                category = "business_scope_not_scheduled"
                actions.append(_action("confirm_scheduling_scope", ["SFXYPK", "LLXS", "ZXS"],
                    {"SFXYPK": flag, "LLXS": course.total_hours, "ZXS": course.hours}, sources,
                    [{"field": "SFXYPK", "candidate_value": 1,
                      "condition": "Only if the responsible office confirms this is a timetable task, not an excluded/practical activity."}],
                    ["业务确认是否需要进入本排课系统；若需要，补齐理论/实践任务口径、总学时、教师责任周及场地要求。"],
                    ["先判断应排范围，不能把274个SFXYPK=0自动计入欠排或直接改为1。"], status="needs_business_confirmation"))
            elif course.hours == 0:
                category = "zero_weekly_hours_positive_total"
                actions.append(_action("confirm_one_off_teaching_weeks", ["ZXS", "LLXS", "SKZCDM"],
                    {"ZXS": course.hours, "LLXS": course.total_hours, "weeks": sorted(course.weeks)}, sources,
                    [{"field": "weekly_load_by_week", "candidate_value": {
                        "weekly_hours": course.total_hours, "week_count": 1, "assigned_weeks": None,
                        "eligible_week_domain": sorted(course.weeks)},
                      "condition": "One-session candidate only; confirm which actual week and whether this is a theory/practice event."}],
                    ["总学时非零，不能免排；确认一次性授课的具体周及教师/班级/场地。确认后调用 preview_data_corrections 的 set_one_off_week 参数，并携带相同 corrections 试排。"],
                    ["不能在所有候选教学周都重复安排总学时；最终合计必须等于已确认LLXS。"], status="needs_confirmation"))
            else:
                category = "high_weekly_load"
                actions.append(_action("configure_concentrated_teaching", ["ZXS", "block_template", "weekly_load_by_week"],
                    {"ZXS": course.hours, "LLXS": course.total_hours, "weeks": sorted(course.weeks), "max_block": course.max_block}, sources,
                    [{"field": "solver_weekly_load_support", "candidate_value": {
                        "required_weekly_load": course.hours, "configured_weekly_limit": config.max_weekly_hours,
                        "supported_weekly_limit": 77, "configured_daily_block_limit": config.max_blocks_per_day,
                        "minimum_blocks_at_declared_max_block": math.ceil(course.hours/course.max_block),
                        "needs_multiple_blocks_per_day_if_workdays": math.ceil(course.hours/course.max_block) > len(config.days)},
                      "condition": "Use explicit max_weekly_hours and max_blocks_per_day parameters; preserve total hours/weeks and do not truncate or round the load."}],
                    ["核实集中教学计划并传入max_weekly_hours、max_blocks_per_day；固定占用补排可设置filter_fixed_occupancy=true。虚班还需确认母课合计总学时。"],
                    ["使用extend_repair_proposal保留现有成功集，完整检查同日课块及资源冲突；不能简单将ZXS截为8。"], status="needs_explicit_parameters"))
                if abs(course.hours-round(course.hours)) > 1e-5 or (course.total_hours is not None and abs(course.hours*len(course.weeks)-course.total_hours)>1e-5):
                    actions.append(_hours_action(course, sources, status))
        else:
            category = status
            if status == "source_hours_inconsistent" or (status == "unsupported" and
                    (not math.isfinite(course.hours) or abs(course.hours-round(course.hours)) > 1e-5
                     or not 1 <= course.hours <= config.max_weekly_hours)):
                actions.append(_hours_action(course, sources, status))
            elif status == "unsupported":
                actions.append(_action("support_reported_model_condition", ["model_condition"],
                    {"detail": failed[cid].get("diagnosis", {}).get("detail", "")}, sources, [],
                    ["按该诊断实现缺失的耦合/连排模型；不能把非课时模型问题当作改LLXS的理由。"],
                    ["保持原声明的课时、教学周、组绑定及资源约束，增加对应回归用例。"], status="needs_implementation"))
            elif status == "data_issue" and any(issue.startswith("coverage:") for issue in course.issues):
                original = _first_values(sources["original"])
                actions.append(_action("restore_verified_class_coverage", ["TJBJ", "class_enrollment_mapping"],
                    {"converted_TJBJ": raw.get("TJBJ", " ".join(sorted(course.classes))),
                     "original_行政班": original.get("行政班"), "original_教学班名称": original.get("教学班名称")},
                    {"source_records": sources,
                     "issues": [re.sub(r"teacher '[^']+'", "teacher [redacted]", issue) for issue in course.issues]},
                    [{"field": "TJBJ", "candidate_value": None,
                      "condition": "Use verified full class IDs or an enrollment-derived conflict mapping; never infer a class from a college/time label."}],
                    ["提供该教学班的真实行政班/选课冲突映射；全校、学院名、时段串和空白均不能替代学生覆盖。",
                     "原源字段若同样异常，需要教务或选课表补证，不能把原异常再复制一次。"],
                    ["班级ID通过注册表校验；覆盖全体实际选课学生，再进行批次冲突检查。"] ))
            elif status == "data_issue":
                actions.append(_action("restore_verified_metadata", ["course_metadata"],
                    {"course_id": cid}, {"source_records": sources,
                        "issues": [re.sub(r"teacher '[^']+'", "teacher [redacted]", issue) for issue in course.issues]}, [],
                    ["按实际缺失或矛盾字段提供权威值及来源，不从课程名/相邻行猜教师、周次或容量。"],
                    ["修正后重新加载课程元数据，全部关键字段及教师责任周覆盖检查通过。"] ))
            elif status == "not_found_within_limits":
                actions.append(_resource_action(course, dataset, state, config, supplemental))
            elif status in {"search_limit", "no_time_pattern", "outside_caller_domain"}:
                actions.append(_action("review_search_coverage", ["search_budget", "time_pattern"],
                    {"status": status}, failed[cid].get("diagnosis", {}), [],
                    ["核对当前候选枚举和预算；先在相同已授权规则下增加预算，不据此判断无解。"],
                    ["保留原成功集并复验整个批次。"], status="automatic_check_only", auto=True))
        if course is not None and category != "business_scope_not_scheduled":
            room_action = _room_action(course, dataset, config, sources)
            if room_action:
                actions.append(room_action)
            original = _first_values(sources["original"])
            actual = _number(original.get("实际人数"))
            if actual is not None and actual > course.capacity:
                actions.append(_action("verify_enrollment_above_capacity_basis", ["KRL", "实际人数"],
                    {"KRL": course.capacity, "source_actual_enrollment": actual}, sources["original"],
                    [{"field": "KRL", "candidate_value": actual, "condition": "Confirm current frozen enrollment; a lower limit must not conceal already enrolled students."}],
                    ["人数已超过当前容量依据，需冻结选课人数及确认是否分班。"],
                    ["所有候选教室容量满足确认人数；不得因上限较小而忽略实际超额。"] ))
        if not actions:
            actions.append(_action("review_unhandled_status", ["status"], {"status": status}, failed.get(cid, {}), [],
                ["补充该状态对应的明确诊断。"], ["所有关键约束均有证据后重试。"] ))
        plans.append({"course_id": cid, "base_course_id": cid.split("@", 1)[0],
            "course_name": course.name if course else skip_values.get("课程名称", ""),
            "origin": "remaining_proposal" if cid in failed else "silently_skipped",
            "status": status, "category": category, "actions": actions})
    return {"schema_version": 1, "mode": "action_plan_only", "source_hashes": dict(dataset.source_hashes),
        "supplemental_source_hashes": supplemental.get("source_hashes", {}),
        "summary": {"remaining_proposal_courses": len(failed), "silently_skipped_courses": len(skipped),
                    "explicitly_promoted_skipped_courses": len(promoted),
                    "total_exact_course_ids": len(plans), "categories": dict(Counter(x["category"] for x in plans))},
        "course_plans": plans,
        "interpretation": "Plans are explicit candidates and evidence requests, not data mutations or permissions. Each corrected group must be jointly revalidated against the current proposal. No enrollment, teacher, class, or calendar-week assignment is fabricated."}
