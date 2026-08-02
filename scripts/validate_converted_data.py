"""
校验「智能排课基础数据/提取的基础数据表_converted」四张表是否符合排课程序要求。

用法：python scripts/validate_converted_data.py [--dir 目录] [--weeks 20 --days 7 --periods 11]
退出码：0=通过（允许有警告），1=存在错误。

校验内容：
  1. 四个文件存在；
  2. 课程表/教室表为两行表头（第1行中文、第2行英文字段名、第3行起数据），
     英文字段名与 Basic_Data 的 Course/Classroom 数据类一致
     （课程表缺列报错、多列忽略；教室表必须恰好等于字段集，多列也报错）；
  3. 教师表（JSH/XM）、班级表（BJMC）为单行英文表头；
  4. 能用 Basic_Data 的 load_* 成功加载；
  5. 内容抽查：SKZCDM 位串长度、ZXS 范围、允许排课教室数（SFYXPK，数字/文本均可）、
     课程 JSH 在教师表中的覆盖率、指定教室代码 JASDM 在教室表中是否存在。
"""
from __future__ import annotations

import argparse
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from openpyxl import load_workbook  # noqa: E402
import pandas as pd  # noqa: E402

COURSE_COLS = [
    "KCH", "KCM", "YXMC", "KCLB", "SKXQ", "JXBID", "KRL", "SKZCDM", "ZXS",
    "JSH", "XM", "RWJSZCDM", "SFXYPK", "SFXYJAS", "JASLXDM", "JASLXMC",
    "JXLDM", "JASDM", "Prefer_Time", "unavailable_Time", "LSJASDM", "TJBJ",
]
CLASSROOM_COLS = [
    "JASDM", "JASMC", "JASLX", "JXLDM", "JXLMC", "LC", "XXXQDM", "MC",
    "SKZWS", "KSZWS", "SFYXPK", "SFYXKS", "SFYXJY",
]

errors: list[str] = []
warnings: list[str] = []


def err(msg: str):
    errors.append(msg)
    print(f"  [错误] {msg}")


def warn(msg: str):
    warnings.append(msg)
    print(f"  [警告] {msg}")


def ok(msg: str):
    print(f"  [通过] {msg}")


def row2_headers(path: str) -> list[str]:
    ws = load_workbook(path, read_only=True).active
    headers = [c.value for c in ws[2]]
    return [h for h in headers if h is not None]


def export_bad_zxs_courses(
    course_excel: str,
    out_path: str = os.path.join("排课结果", "周学时异常课程_待教务确认.xlsx"),
) -> str:
    """导出 ZXS 不在 (0,8] 的课程全部原始字段，附说明页供教务老师逐条确认。"""
    df = pd.read_excel(course_excel, header=1)
    z = pd.to_numeric(df["ZXS"], errors="coerce")
    bad = df[~((z > 0) & (z <= 8))].copy()
    bad.insert(0, "教务处理意见", "")

    dist = bad["ZXS"].value_counts(dropna=False)
    dist_txt = "；".join(f"ZXS={k}: {v}行" for k, v in dist.head(8).items())
    notes = pd.DataFrame({"说明": [
        "排课系统仅支持周学时(ZXS)在 1–8 之间的课程，下列课程因 ZXS 超出范围被系统自动跳过：",
        "**不参与排课，也不会出现在排课结果或失败清单中**。请逐条确认并在明细表「教务处理意见」列填写：",
        "",
        "1. 这门课是否需要系统排课？若不需要（如 MOOC、实践环节、毕业设计等），请填「无需排课」；",
        "2. 若需要排课且 ZXS=0 或为空：请填写实际周学时（1–8）；",
        "3. 若 ZXS>8（如 16/60）：是否把学期总学时误填为周学时？请填写正确周学时，",
        "   或说明特殊排课规则（如分组教学、集中数周上课等），由系统另行处理；",
        "",
        f"本次共 {len(bad)} 行，ZXS 取值分布：{dist_txt}",
        "确认后请将本文件返还，由数据整理脚本回填课程表。",
    ]})
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        notes.to_excel(writer, sheet_name="说明_请教务老师确认", index=False)
        bad.to_excel(writer, sheet_name="周学时异常课程明细", index=False)
    return out_path


def check_two_row_table(path: str, required: list[str], exact: bool, label: str):
    print(f"\n--- {label}: {os.path.basename(path)} ---")
    headers = row2_headers(path)
    missing = [c for c in required if c not in headers]
    extra = [c for c in headers if c not in required]
    if missing:
        err(f"第2行英文表头缺少字段: {missing}")
    if extra:
        if exact:
            err(f"存在多余字段（加载会直接报 TypeError，必须删除）: {extra}")
        else:
            # 排课引擎只读 Course 数据类声明的字段，其余列（如 PKYQMS 排课要求原文、
            # course_schedule_pipeline 的中间列）由上游模块使用，保留即可
            print(f"  [说明] 以下字段排课引擎不读取，但上游模块（需求解析/数据准备）会用，保留无碍: {extra}")
    if not missing and not (exact and extra):
        ok(f"表头字段完整（{len(headers)} 列）")


def main():
    parser = argparse.ArgumentParser(description="校验排课基础数据四张表")
    parser.add_argument("--dir", default=os.path.join("智能排课基础数据", "提取的基础数据表_converted"))
    parser.add_argument("--weeks", type=int, default=20)
    parser.add_argument("--days", type=int, default=7)
    parser.add_argument("--periods", type=int, default=11)
    args = parser.parse_args()

    paths = {
        "课程表": os.path.join(args.dir, "课程表.xlsx"),
        "教室表": os.path.join(args.dir, "教室表.xlsx"),
        "教师表": os.path.join(args.dir, "教师表.xlsx"),
        "班级表": os.path.join(args.dir, "班级表.xlsx"),
    }
    for label, p in paths.items():
        if not os.path.isfile(p):
            err(f"{label}不存在: {p}")
    if errors:
        print(f"\n校验失败：{len(errors)} 个错误")
        sys.exit(1)

    check_two_row_table(paths["课程表"], COURSE_COLS, exact=False, label="课程表")
    check_two_row_table(paths["教室表"], CLASSROOM_COLS, exact=True, label="教室表")

    print("\n--- 教师表 / 班级表 ---")
    df_t = pd.read_excel(paths["教师表"])
    for col in ("JSH", "XM"):
        if col not in df_t.columns:
            err(f"教师表缺少列 {col}（单行英文表头）")
    df_b = pd.read_excel(paths["班级表"])
    if "BJMC" not in df_b.columns:
        err("班级表缺少列 BJMC（单行英文表头）")
    if not errors:
        ok(f"教师表 {len(df_t)} 行 / 班级表 {len(df_b)} 行")

    if errors:
        print(f"\n校验失败：{len(errors)} 个错误（表头不符，跳过加载测试）")
        sys.exit(1)

    print("\n--- 用 Basic_Data 加载器实际加载 ---")
    from Basic_Data import load_courses, load_classrooms, load_teachers, load_classes
    try:
        courses = load_courses(paths["课程表"], weeks=args.weeks, days=args.days, periods=args.periods)
        classrooms = load_classrooms(paths["教室表"], weeks=args.weeks, days=args.days, periods=args.periods)
        teachers = load_teachers(paths["教师表"], weeks=args.weeks, days=args.days, periods=args.periods)
        load_classes(paths["班级表"], weeks=args.weeks, days=args.days, periods=args.periods)
        ok(f"加载成功：课程行 {len(courses)} / 教室 {len(classrooms)} / 教师 {len(teachers)}")
    except Exception as e:
        err(f"加载失败: {type(e).__name__}: {e}")
        print(f"\n校验失败：{len(errors)} 个错误")
        sys.exit(1)

    print("\n--- 内容抽查 ---")
    n_bad_skzcdm = sum(
        1 for c in courses
        if c.SKZCDM is None or not str(c.SKZCDM).strip() or len(str(c.SKZCDM).strip()) < args.weeks
    )
    if n_bad_skzcdm:
        warn(f"{n_bad_skzcdm} 行课程 SKZCDM 为空或长度小于学期周数 {args.weeks}")
    else:
        ok("SKZCDM 周次位串长度正常")

    # 与排课引擎 pre_selection 的条件完全一致：0 < ZXS <= 8（数值比较）
    n_bad_zxs = 0
    for c in courses:
        try:
            z = float(c.ZXS)
        except (TypeError, ValueError):
            z = float("nan")
        if not (0 < z <= 8):
            n_bad_zxs += 1
    if n_bad_zxs:
        export = export_bad_zxs_courses(paths["课程表"])
        warn(
            f"{n_bad_zxs} 行课程 ZXS 不在 (0,8] 范围。注意：排课程序会把这些课**静默跳过**"
            f"——既不参与排课，也不会出现在失败清单里。已导出全部课程信息到「{export}」，"
            f"请交教务老师确认：这些课需不需要系统排课？周学时正确值是多少？是否有特殊排课逻辑？"
        )
    else:
        ok("ZXS 周学时范围正常")

    n_pk_rooms = sum(1 for c in classrooms if str(c.SFYXPK).strip() in ("1", "1.0"))
    if n_pk_rooms == 0:
        err("没有任何教室 SFYXPK=1（允许排课），排课必然全失败")
    else:
        ok(f"允许排课教室 {n_pk_rooms} 间")

    jsh_set = {str(j).strip() for j in df_t["JSH"].dropna()}
    course_jsh = {str(c.JSH).strip() for c in courses if c.JSH is not None and str(c.JSH).strip()}
    missing_jsh = course_jsh - jsh_set
    if missing_jsh:
        ratio = len(missing_jsh) / max(len(course_jsh), 1) * 100
        warn(f"{len(missing_jsh)} 个课程教师号不在教师表中（占 {ratio:.1f}%），这些课会因教师缺失排不上")
    else:
        ok("课程教师号在教师表中全覆盖")

    room_set = {str(c.JASDM).strip() for c in classrooms}
    bad_rooms = {
        str(c.JASDM).strip() for c in courses
        if c.JASDM is not None and str(c.JASDM).strip() and str(c.JASDM).strip() not in room_set
    }
    if bad_rooms:
        warn(f"{len(bad_rooms)} 个课程指定教室代码不在教室表中，例如: {sorted(bad_rooms)[:5]}")
    else:
        ok("课程指定教室代码均存在于教室表")

    print("\n" + "=" * 50)
    if errors:
        print(f"校验失败：{len(errors)} 个错误，{len(warnings)} 个警告")
        sys.exit(1)
    print(f"校验通过：0 个错误，{len(warnings)} 个警告")


if __name__ == "__main__":
    main()
