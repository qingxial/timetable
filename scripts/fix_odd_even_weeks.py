# -*- coding: utf-8 -*-
"""修复 converted 课程表里被「X单/X双」格式拖坏的 SKZCDM 字段。
- 源：排课要求和基础数据_v1.0.xlsx（教学任务 sheet）
- 目标：智能排课基础数据/提取的基础数据表_converted/课程表.xlsx
匹配键：JXBID（即课程序号），用源「起始周/结束周」重算 SKZCDM 与 SKZCMC。
RWJSZCDM 若与 SKZCDM 完全一致也同步重算。原文件先备份为 .bak.YYYYMMDD_HHMM。
默认 --dry-run 只打印对比；加 --apply 才真改文件。
"""
import argparse, os, re, shutil, datetime
import pandas as pd
from openpyxl import load_workbook

SRC_FILE = "排课要求和基础数据_v1.0.xlsx"
DST_FILE = "智能排课基础数据/提取的基础数据表_converted/课程表.xlsx"
TOTAL_WEEKS = 23   # SKZCDM 位串长度，与现有 converted 一致

def parse_end_week(end_raw):
    """'17单' -> (17, 'odd'); '16双' -> (16, 'even'); '9' -> (9, None)"""
    s = str(end_raw).strip()
    m = re.match(r'^\s*(\d+)\s*([单双])?\s*$', s)
    if not m:
        return None, None
    end = int(m.group(1))
    parity = {'单': 'odd', '双': 'even'}.get(m.group(2))
    return end, parity

def build_skzcdm(start, end, parity, total=TOTAL_WEEKS):
    """生成 23 位 0/1 字符串。周次从 1 开始；位索引 [w-1]。"""
    bits = ['0'] * total
    for w in range(int(start), int(end) + 1):
        if w < 1 or w > total: continue
        if parity == 'odd'  and w % 2 == 0: continue
        if parity == 'even' and w % 2 == 1: continue
        bits[w-1] = '1'
    return ''.join(bits)

def build_skzcmc(skzcdm):
    """从位串还原出 '2,3,4...' 这种逗号串。"""
    return ','.join(str(i+1) for i, b in enumerate(skzcdm) if b == '1')

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="真改文件，否则只 dry-run")
    args = ap.parse_args()

    # 1) 读源教学任务，建 JXBID -> (起始周, 结束周原始)
    src = pd.read_excel(SRC_FILE, sheet_name="教学任务")
    id_col = "课程序号"
    src_map = {}
    for _, r in src.iterrows():
        jxbid = str(r[id_col]).strip()
        src_map[jxbid] = (r["起始周"], r["结束周"])

    # 2) 读 converted 课程表（保留原始 2 行表头）
    raw = pd.read_excel(DST_FILE, header=None)
    header_cn, header_en = raw.iloc[0], raw.iloc[1]
    data = raw.iloc[2:].copy().reset_index(drop=True)
    data.columns = header_en

    skz_col = "SKZCDM"
    smc_col = "SKZCMC"
    rwj_col = "RWJSZCDM"
    jxb_col = "JXBID"

    mask_zero = data[skz_col].astype(str).str.strip() == "0" * TOTAL_WEEKS
    print(f"converted 总行数: {len(data)}, SKZCDM 全 0 行数: {mask_zero.sum()}")

    fixed = 0; samples = []
    for i in data.index[mask_zero]:
        jxbid = str(data.at[i, jxb_col]).strip()
        if jxbid not in src_map:
            continue
        start, end_raw = src_map[jxbid]
        end, parity = parse_end_week(end_raw)
        if end is None or parity is None:
            continue  # 无法解析或无单双标记 → 跳过
        try:
            start = int(start)
        except (TypeError, ValueError):
            continue
        new_skz = build_skzcdm(start, end, parity)
        if new_skz == "0" * TOTAL_WEEKS:
            continue
        old_rwj = str(data.at[i, rwj_col]).strip() if rwj_col in data.columns else ""
        rwj_sync = (old_rwj == "0" * TOTAL_WEEKS)

        if len(samples) < 6:
            samples.append((jxbid, str(data.at[i,'KCM']), start, end_raw, parity, new_skz, build_skzcmc(new_skz)))

        if args.apply:
            data.at[i, skz_col] = new_skz
            data.at[i, smc_col] = build_skzcmc(new_skz)
            if rwj_sync:
                data.at[i, rwj_col] = new_skz
        fixed += 1

    print(f"将修复 {fixed} 行")
    print("\n样例（前 6 行）：")
    print(f"{'JXBID':<22} {'课程':<12} 起始 结束  奇偶  SKZCDM(23位)              SKZCMC")
    for s in samples:
        print(f"{s[0]:<22} {s[1]:<12} {s[2]:>4} {str(s[3]):>4}  {s[4]:<5} {s[5]}  {s[6]}")

    if not args.apply:
        print("\n[dry-run] 未写入。加 --apply 真改文件（会自动备份）。")
        return

    # 3) 备份；用 openpyxl 原地改单元格（避免列名重复导致的 concat 报错，保留原表头/列序/格式）
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M")
    bak = DST_FILE + f".bak.{ts}"
    shutil.copy2(DST_FILE, bak)
    print(f"\n已备份原文件 → {bak}")

    wb = load_workbook(DST_FILE)
    ws = wb.active
    # 第 2 行（openpyxl 1-based）= 英文表头
    en_header = [c.value for c in ws[2]]
    def col_idx(name):
        try: return en_header.index(name) + 1   # openpyxl 1-based
        except ValueError: return None
    c_jxb = col_idx("JXBID"); c_skz = col_idx(skz_col); c_smc = col_idx(smc_col); c_rwj = col_idx(rwj_col)
    if not all([c_jxb, c_skz, c_smc, c_rwj]):
        raise RuntimeError(f"找不到必要列：JXBID/{skz_col}/{smc_col}/{rwj_col}")

    written = 0
    # 数据从第 3 行开始
    for row in range(3, ws.max_row + 1):
        cur_skz = str(ws.cell(row, c_skz).value or "").strip()
        if cur_skz != "0" * TOTAL_WEEKS:
            continue
        jxbid = str(ws.cell(row, c_jxb).value or "").strip()
        if jxbid not in src_map: continue
        start, end_raw = src_map[jxbid]
        end, parity = parse_end_week(end_raw)
        if end is None or parity is None: continue
        try: start = int(start)
        except (TypeError, ValueError): continue
        new_skz = build_skzcdm(start, end, parity)
        if new_skz == "0" * TOTAL_WEEKS: continue
        old_rwj = str(ws.cell(row, c_rwj).value or "").strip()
        ws.cell(row, c_skz).value = new_skz
        ws.cell(row, c_smc).value = build_skzcmc(new_skz)
        if old_rwj == "0" * TOTAL_WEEKS:
            ws.cell(row, c_rwj).value = new_skz
        written += 1
    wb.save(DST_FILE)
    print(f"已写回 → {DST_FILE}  （实际改写 {written} 行）")

if __name__ == "__main__":
    main()
